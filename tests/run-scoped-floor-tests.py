import copy
import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from tools import update_excel_rates as engine

BANDS = [(1, 1, 200), (2, 2, 190), (3, 4, 180), (5, 7, 170), (8, 20, 160), (21, 35, 150)]


class ScopedFloorTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.pricing = json.loads((ROOT / 'pricing-rules.config.example.json').read_text())
        self.floors = [dict(groups=['CDMV', 'CFAV'], zones=['WA1', 'WALO'],
                            startDate='2026-10-15', endDate='2026-11-02', minDays=low, maxDays=high,
                            minimumRatePlnDay=rate, overrideProtected=False) for low, high, rate in BANDS]
        self.config = engine.load_config(ROOT / 'excel-rate-update.config.example.json')
        self.config.update(pickup_date_expansion={'enabled': False}, baseline_manifest_file='',
                           pricing_rules_file=str(self.root / 'pricing.json'))
        self.write_rules()

    def write_rules(self):
        self.pricing['pricing']['mandatoryScopedFloorsPlnDay'] = self.floors
        (self.root / 'pricing.json').write_text(json.dumps(self.pricing))

    def workbook(self):
        book = openpyxl.Workbook()
        ws = book.active
        ws.title = 'Sheet1'
        ws.append(['Rental rates'])
        ws.append([None] * 8 + [1, 2, 3, 5, 8, 21])
        ws.append([None] * 8 + [1, 2, 4, 7, 20, 35])
        ws.append(['Group', 'Description', 'Rate code', 'Zone', 'Booking start date', 'Booking end date', 'Pickup start date', 'Pickup end date'] + ['Per day'] * 6)
        for pickup in ['2026-10-14', '2026-10-15', '2026-10-31', '2026-11-02', '2026-11-03']:
            for zone in ['WA1', 'WALO', 'WR1']:
                for group in ['CDMV', 'CFAV', 'SWAV', 'EDAV']:
                    high = group == 'CFAV' and zone == 'WA1' and pickup == '2026-10-15'
                    ws.append([group, None, None, zone] + [pickup] * 4 + [500 if high else 80] * 6)
        source = self.root / 'source.xlsx'
        book.save(source)
        book.close()
        return source

    def test_real_excel_updates_scopes_frozen_classes_dates_all_bands_and_dry_run(self):
        source = self.workbook()
        before = source.read_bytes()
        recs = self.root / 'recs.json'
        recs.write_text('{"decisions": []}')
        preview = engine.apply_updates(source, recs, None, self.config, None, True)
        output, imported = self.root / 'review.xlsx', self.root / 'import.xlsx'
        summary = engine.apply_updates(source, recs, output, self.config, None, False, import_output_path=imported)
        self.assertEqual(preview['changes'], summary['changes'])
        self.assertTrue(summary['changes'])
        self.assertFalse(any(row['status'] == 'FAIL' for row in summary['validation']))
        for path in [output, imported]:
            book = openpyxl.load_workbook(path)
            for row in book['Sheet1'].iter_rows(min_row=5, values_only=True):
                group, zone, pickup = row[0], row[3], str(row[6])[:10]
                active = group in ['CDMV', 'CFAV'] and zone in ['WA1', 'WALO'] and pickup in ['2026-10-15', '2026-11-02']
                high = group == 'CFAV' and zone == 'WA1' and pickup == '2026-10-15'
                existing = 500 if high else 110 if pickup == '2026-10-31' else 80
                self.assertEqual(list(row[8:14]), [max(rate, existing) if active else existing for _, _, rate in BANDS], (group, zone, pickup))
            if path == imported:
                self.assertEqual(book.sheetnames, ['Sheet1'])
                self.assertTrue(all(c.comment is None for row in book['Sheet1'] for c in row))
            book.close()
        self.assertEqual(source.read_bytes(), before)

    def test_public_resolvers_match_js_for_scopes_holidays_and_overlap(self):
        self.floors.append({**self.floors[1], 'minimumRatePlnDay': 100})
        self.write_rules()
        cases = [('WA1', 'CFAV', 2, '2026-10-15', 190), ('WR1', 'CFAV', 2, '2026-10-15', None),
                 ('WA1', 'SWAV', 2, '2026-10-15', None), ('WA1', 'CFAV', 2, '2026-10-31', None),
                 ('WA1', 'CFAV', 35, '2026-11-02', 150), ('WA1', 'CFAV', 2, '2026-11-03', None)]
        periods = self.config['protected_rate_periods']
        script = "const e=require('./src/pricingRules');const x=JSON.parse(process.argv[1]);console.log(JSON.stringify(x.cases.map(([z,g,d,p])=>e.resolveMandatoryScopedFloorPlnDay(z,d,{mandatoryScopedFloorsPlnDay:x.floors},g,p,x.periods))));"
        js = json.loads(subprocess.check_output(['node', '-e', script, json.dumps(dict(cases=cases, floors=self.floors, periods=periods))], cwd=ROOT, text=True))
        for case, actual in zip(cases, js):
            zone, group, days, pickup, expected = case
            self.assertEqual(engine.get_mandatory_scoped_floor(zone, group, days, days, self.config, pickup), expected)
            self.assertEqual(actual, expected or 0)

    def test_existing_holiday_and_premium_correction_prevail_over_lower_scoped_floor(self):
        self.floors = [{**self.floors[1], 'groups': ['FVMD'], 'startDate': '2026-12-15', 'endDate': '2027-01-10'}]
        self.write_rules()
        self.assertEqual(engine.get_mandatory_floor('WA1', 'FVMD', 2, 2, self.config, '2026-12-25')[0], 700)
        self.assertIsNone(engine.get_mandatory_scoped_floor('WA1', 'FVMD', 2, 2, self.config, '2027-01-10'))
        self.floors[0]['overrideProtected'] = True
        self.write_rules()
        self.assertEqual(engine.get_mandatory_floor('WA1', 'FVMD', 2, 2, self.config, '2027-01-10')[0], 190)
        self.floors = [{**self.floors[0], 'groups': ['EDAV'], 'startDate': '2026-10-15', 'endDate': '2026-11-02', 'minimumRatePlnDay': 60.5}]
        self.write_rules()
        self.assertEqual(engine.get_mandatory_correction_floor('WA1', 'EDAV', 2, 2, self.config, '2026-10-15')[0], 61)

    def test_invalid_schema_and_unknown_scopes_fail_closed(self):
        original = copy.deepcopy(self.floors[0])
        for patch in [{'groups': ['ZZZZ']}, {'groups': ['*']}, {'zones': ['FAKE']}, {'groups': []},
                      {'startDate': '2026-02-30'}, {'maxDays': 3}, {'minimumRatePlnDay': True},
                      {'overrideProtected': 'false'}, {'kind': 'city'}]:
            self.floors = [{**original, **patch}]
            self.write_rules()
            with self.subTest(patch=patch), self.assertRaisesRegex(ValueError, 'mandatoryScopedFloorsPlnDay'):
                engine.get_mandatory_floor('WA1', 'CFAV', 1, 1, self.config, '2026-10-15')

    def test_missing_duration_columns_and_tampered_final_rate_are_detected(self):
        source = self.workbook()
        book = openpyxl.load_workbook(source)
        ws = book['Sheet1']
        columns = engine.get_duration_columns(ws, self.config)
        with self.assertRaisesRegex(ValueError, 'Missing Excel duration columns'):
            engine.enforce_mandatory_floors(ws, self.config, {day: col for day, col in columns.items() if day != 35}, True, [])
        changes = engine.enforce_mandatory_floors(ws, self.config, columns, False, [])
        self.assertTrue(changes)
        checks = engine.build_validation_rows(ws, self.config, columns, changes, [])
        self.assertFalse(any(row[1] == 'FAIL' for row in checks))
        change = next(change for change in changes if change['group'] == 'CFAV' and change['pickup_date'] == '2026-11-02')
        ws[change['cell']].value = 1
        checks = engine.build_validation_rows(ws, self.config, columns, changes, [])
        self.assertGreater(next(row[2] for row in checks if row[0] == 'Stawki ponizej obowiazkowego minimum klasy'), 0)
        book.close()

    def test_real_recommendation_parity_cannot_undo_a_selected_premium_floor(self):
        self.floors = [{**self.floors[1], 'groups': ['EDAV'], 'minimumRatePlnDay': 250}]
        self.write_rules()
        source = self.workbook()
        recs = self.root / 'recs.json'
        recs.write_text(json.dumps({'decisions': [{
            'action': 'increase', 'recommendation_type': 'top1_gap', 'location': 'Warsaw West Train Station',
            'start_date': '2026-11-02', 'rental_days': 2, 'suggested_rate_pln_day': 200,
            'maximum_import_rate_pln_day': 200, 'site_cap_rate_pln_day': 200,
            'benchmark_rate_pln_day': 201, 'mm_rate_pln_day': 80, 'data_quality_status': 'ok',
            'target_rank': 1, 'scenario_id': 'parity-floor',
        }]}))
        result = engine.apply_updates(source, recs, None, self.config, None, True)
        self.assertFalse(any(row['status'] == 'FAIL' for row in result['validation']))
        scoped = [change for change in result['changes'] if change['group'] == 'EDAV'
                  and change['zone'] == 'WA1' and change['pickup_date'] == '2026-11-02' and change['duration_min_days'] == 2]
        self.assertTrue(scoped)
        self.assertGreaterEqual(scoped[-1]['new_rate'], 250)


if __name__ == '__main__':
    unittest.main()
