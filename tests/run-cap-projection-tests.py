import json
import sys
import tempfile
import unittest
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from tools.update_excel_rates import apply_updates, merge_config  # noqa: E402


GROUPS = ["CDMV", "CGAV", "CWAV", "CWMR", "EDAV", "EDMV"]


class CapProjectionTests(unittest.TestCase):
    def test_airport_family_minimum_keeps_compatible_city_cap_active(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "baseline.xlsx"
            output = root / "review.xlsx"
            imported = root / "import.xlsx"
            pickup = "2026-10-26"
            book = openpyxl.Workbook()
            sheet = book.active
            sheet.title = "Sheet1"
            sheet.append(["Rental rates for packages: INCLUSIVE POA"])
            sheet.append([None] * 8 + [1, 2, 3, 5, 8, 21])
            sheet.append([None] * 8 + [1, 2, 4, 7, 20, 35])
            sheet.append(["Group", "Description", "Rate code", "Zone", "Booking start date",
                          "Booking end date", "Pickup start date", "Pickup end date"] + ["Per day"] * 6)
            for zone, rate in (("WA1", 60), ("WALO", 40)):
                for group in GROUPS:
                    premium = int(group in {"EDAV", "EDMV"})
                    sheet.append([group, None, None, zone] + [pickup] * 4 + [rate + premium] * 6)
            book.save(source)
            book.close()
            original = source.read_bytes()

            rules = json.loads((ROOT / "pricing-rules.config.example.json").read_text(encoding="utf-8"))
            rules["pricing"]["mandatoryGroupFloorsPlnDay"]["EDAV"] = [
                {"minDays": 2, "maxDays": 2, "minimumRatePlnDay": 100},
            ]
            rules_path = root / "pricing.json"
            rules_path.write_text(json.dumps(rules), encoding="utf-8")
            config = merge_config({
                "pricing_rules_file": str(rules_path),
                "location_zones": {"City": ["WA1"], "Airport": ["WALO"]},
                "city_zone_airport_zones": {"WA1": ["WALO"]},
                "zone_location_kinds": {"WA1": "city", "WALO": "airport"},
                "pickup_date_expansion": {"enabled": False},
                "minimum_rates": {"global_min_pln_day": 0, "bands": [], "scoped_overrides": []},
            })
            decisions = [{
                "action": "increase", "recommendation_type": "top1_gap",
                "location": location, "start_date": pickup, "rental_days": 2,
                "suggested_rate_pln_day": target, "maximum_import_rate_pln_day": target,
                "site_cap_rate_pln_day": target, "benchmark_rate_pln_day": target + 1,
                "mm_rate_pln_day": baseline, "data_quality_status": "ok", "target_rank": 1,
                "scenario_id": f"{location}-{pickup}-2",
            } for location, target, baseline in (("City", 150, 60), ("Airport", 40, 40))]
            recommendations = root / "recommendations.json"
            recommendations.write_text(json.dumps({"decisions": decisions}), encoding="utf-8")

            summary = apply_updates(source, recommendations, output, config, None, False,
                                    import_output_path=imported)
            for path in (output, imported):
                result = openpyxl.load_workbook(path)
                try:
                    prices = {(row[0], row[3]): row[9]
                              for row in result["Sheet1"].iter_rows(min_row=5, values_only=True)}
                    for group in GROUPS:
                        premium = int(group in {"EDAV", "EDMV"})
                        self.assertEqual(prices[group, "WALO"], 100 + premium)
                        # Airport parity raises the base to 100, so the compatible city cap is 130.
                        self.assertEqual(prices[group, "WA1"], 130 + premium,
                                         f"{path.name}/{group}: city must respect projected airport cap")
                        self.assertLessEqual(prices[group, "WA1"], prices[group, "WALO"] * 1.3)
                finally:
                    result.close()
            self.assertEqual(summary["city_top1_airport_cap_ignored_count"], 0)
            self.assertEqual(summary["city_top1_airport_cap_violation_count"], 0)
            self.assertEqual(source.read_bytes(), original)


if __name__ == "__main__":
    unittest.main(verbosity=2)
