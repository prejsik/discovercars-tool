import json
import posixpath
import shutil
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from xml.etree import ElementTree
from zipfile import ZipFile

import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import PatternFill

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "tests"))

from tools.update_excel_rates import (  # noqa: E402
    apply_updates, build_validation_rows, get_duration_columns, get_review_status,
    load_baseline_confirmation, get_import_row_limit, validate_import_row_limit, parse_date_value,
    load_config, get_minimum_rate, get_mandatory_floor,
    save_import_ready_workbook,
)
from tools import update_excel_rates as excel_updater  # noqa: E402
from excel_test_fixtures import (  # noqa: E402
    FIXTURE_PRICING_RULES_PATH, FIXTURE_REGISTRY_PATH, load_fixture_config,
    load_fixture_pricing_rules, merge_fixture_config,
)


ALLOWED = ["CDMV", "CGAV", "CWAV", "CWMR", "EDAV", "EDMV"]
FROZEN = ["CFAV", "PDAH", "PDAV", "FVMD", "SWAV", "MDMR"]


class PricingPolicyTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        legacy_rules = load_fixture_pricing_rules()
        legacy_rules["pricing"].pop("locationTypeFloorsPlnDay", None)
        legacy_path = self.root / "legacy-rules.json"
        legacy_path.write_text(json.dumps(legacy_rules), encoding="utf-8")
        self.config = merge_fixture_config({
            "pricing_rules_file": str(legacy_path),
            "location_zones": {"City": ["WA1"]},
            "city_zone_airport_zones": {"WA1": ["WALO"]},
            "protected_rate_periods": [
                {"start_date": "2026-10-31", "end_date": "2026-11-01"},
                {"start_date": "2026-12-15", "end_date": "2027-01-10"},
            ],
            "minimum_rates": {"global_min_pln_day": 50, "bands": [], "scoped_overrides": [
                {"zones": ["*"], "groups": ALLOWED, "min_days": 8, "max_days": 20, "min_pln_day": 40},
            ]},
            # Isolate legacy market/zone fixtures; group minima have dedicated tests below.
            "_mandatory_group_floors": {},
        })

    def pdah_floor_config(self):
        return {key: value for key, value in self.config.items() if key != "_mandatory_group_floors"}

    def test_all_classes_october_holiday_floor_boundaries_and_higher_rates(self):
        dates = ("2026-10-29", "2026-10-30", "2026-10-31", "2026-11-01", "2026-11-02")
        source = self.workbook(pickups=dates, airport_rate=80, city_rate=180)
        original = source.read_bytes()
        _, output, import_output = self.run_updates([], source, self.pdah_floor_config())
        for path in (output, import_output):
            for col in range(9, 15):
                for (group, zone, pickup), price in self.rates(path, col).items():
                    old = (80 if zone == "WALO" else 180) + (group in {"EDAV", "EDMV"})
                    active = "2026-10-30" <= pickup <= "2026-11-01"
                    minimum = 170 if group == "PDAH" and col == 14 and pickup not in {"2026-10-31", "2026-11-01"} else 0
                    self.assertEqual(price, max(old, minimum, 110 if active else 0), (group, zone, pickup, col))
        self.assertEqual(source.read_bytes(), original)

    def test_fvmd_holiday_minima_boundaries_all_bands_and_preserved_higher_rates(self):
        dates = ("2026-12-14", "2026-12-15", "2026-12-25", "2027-01-05", "2027-01-06", "2027-01-10")
        source = self.workbook(pickups=dates, airport_rate=1200, city_rate=1200)
        book = openpyxl.load_workbook(source)
        for row in book["Sheet1"].iter_rows(min_row=5):
            if row[0].value == "FVMD" and row[3].value == "WA1":
                for cell in row[8:14]:
                    cell.value = 100
        book["Sheet1"].append(["FVMD", None, None, "SZLO"] + ["2026-12-25"] * 4 + [100] * 6)
        book.save(source)
        book.close()
        original = source.read_bytes()
        summary, output, import_output = self.run_updates([], source, self.pdah_floor_config())
        minima = (1000, 700, 600, 500, 500, 400)
        for path in (output, import_output):
            for col, minimum in enumerate(minima, 9):
                prices = self.rates(path, col)
                self.assertFalse(any(key[:2] == ("FVMD", "SZLO") for key in prices))
                for key, price in prices.items():
                    active = key[0] == "FVMD" and key[1] == "WA1" and "2026-12-15" <= key[2] <= "2027-01-05"
                    old = 100 if key[0] == "FVMD" and key[1] == "WA1" else 1200 + (key[0] in {"EDAV", "EDMV"})
                    self.assertEqual(price, minimum if active else old, (key, col))
        self.assertEqual(summary["mandatory_group_floor_change_count"], 18)
        self.assertFalse(any(row["status"] == "FAIL" for row in summary["validation"]))
        self.assertEqual(source.read_bytes(), original)
        book = openpyxl.load_workbook(output)
        legend = " ".join(str(c.value or "") for row in book["Changed Positions"].iter_rows() for c in row)
        self.assertIn("FVMD", legend)
        self.assertIn("2026-12-15", legend)
        self.assertIn("2027-01-05", legend)
        book.close()

    def test_fvmd_holiday_dry_run_and_validator_reject_breach_or_unauthorized_change(self):
        source = self.workbook(pickups=("2026-12-15",), airport_rate=100, city_rate=100)
        config = self.pdah_floor_config()
        recommendations = self.root / "empty.json"
        recommendations.write_text('{"decisions": []}', encoding="utf-8")
        preview = apply_updates(source, recommendations, None, config, None, True)
        summary, output, _ = self.run_updates([], source, config)
        self.assertEqual(preview["changes"], summary["changes"])
        self.assertEqual(summary["mandatory_group_floor_change_count"], 12)
        book = openpyxl.load_workbook(output)
        sheet = book["Sheet1"]
        row = next(r[0].row for r in sheet.iter_rows(min_row=5) if r[0].value == "FVMD")
        sheet.cell(row, 9).value = 999
        checks = build_validation_rows(sheet, config, get_duration_columns(sheet, config), [], [])
        self.assertEqual(next(c for c in checks if c[0] == "Stawki ponizej obowiazkowego minimum klasy")[2], 1)
        fake = {**summary["changes"][0], "new_rate": 1001}
        checks = build_validation_rows(sheet, config, get_duration_columns(sheet, config), [fake], [])
        self.assertEqual(next(c for c in checks if c[0] == "Zmienione stawki w chronionych datach")[2], 1)
        book.close()

    def test_fvmd_holiday_invalid_scope_fails_closed(self):
        original = load_fixture_pricing_rules()
        for override in ({"start_date": "bad"}, {"end_date": "2026-12-14"},
                         {"override_protected_dates": "true"}, {"end_date": None}):
            rules = json.loads(json.dumps(original))
            rules["pricing"]["mandatoryGroupFloorsPlnDay"]["FVMD"][0].update(override)
            path = self.root / "bad-holiday-scope.json"
            path.write_text(json.dumps(rules), encoding="utf-8")
            config = {**self.pdah_floor_config(), "pricing_rules_file": str(path)}
            with self.subTest(override=override), self.assertRaisesRegex(ValueError, "date scope"):
                self.run_updates([], self.workbook(), config)

    def test_pdah_long_floor_covers_all_zones_dates_and_both_outputs(self):
        source = self.workbook(pickups=("2026-10-01", "2028-02-10"), airport_rate=450, city_rate=450)
        book = openpyxl.load_workbook(source)
        sheet = book["Sheet1"]
        sheet.append(["PDAH", None, None, "SZLO"] + ["2026-10-01"] * 4 + [450] * 5 + [100])
        sheet.append(["CFAV", None, None, "SZLO"] + ["2026-10-01"] * 4 + [450] * 5 + [100])
        pdah_rates = iter((100, 169.99, 170, 220))
        for row in sheet.iter_rows(min_row=5):
            if row[0].value == "PDAH" and row[3].value != "SZLO":
                row[13].value = next(pdah_rates)
                sheet.row_dimensions[row[0].row].hidden = True
        book.save(source)
        book.close()
        source_bytes = source.read_bytes()
        before = {column: self.rates(source, column) for column in range(9, 15)}
        summary, output, import_output = self.run_updates([], source, self.pdah_floor_config())
        for path in (output, import_output):
            for column, old_prices in before.items():
                prices = self.rates(path, column)
                for key, old_rate in old_prices.items():
                    expected = max(old_rate, 170) if key[0] == "PDAH" and column == 14 else old_rate
                    self.assertEqual(prices[key], expected)
        self.assertEqual(summary["mandatory_group_floor_change_count"], 3)
        self.assertEqual(summary["mandatory_group_floor_groups"], ["PDAH"])
        self.assertEqual(summary["mandatory_zone_floor_change_count"], 0)
        self.assertEqual(summary["change_statistics"]["decrease_count"], 0)
        self.assertEqual(summary["recommendation_date_scope"]["date_count"], 0)
        self.assertEqual(source.read_bytes(), source_bytes)
        self.assertFalse(any(row["status"] == "FAIL" for row in summary["validation"]))
        book = openpyxl.load_workbook(output)
        legend = " ".join(str(cell.value or "") for row in book["Changed Positions"].iter_rows() for cell in row)
        self.assertIn("PDAH", legend)
        self.assertIn("21-35: 170 PLN", legend)
        review_rows = list(book["Recommendations Review"].iter_rows(min_row=2, values_only=True))
        self.assertTrue(review_rows)
        for row in review_rows:
            self.assertEqual(row[1], "Gotowe z uwaga")
            self.assertIn("bez celu rankingowego", row[2])
            self.assertEqual(row[14], "Obowiazkowe minimum klasy")
        book.close()

    def test_pdah_long_floor_preserves_protected_dates_and_other_bands(self):
        protected = {"2026-10-31", "2026-11-01", "2026-12-15", "2027-01-10"}
        source = self.workbook(pickups=tuple(sorted(protected | {"2027-01-11"})), airport_rate=500, city_rate=500)
        book = openpyxl.load_workbook(source)
        for row in book["Sheet1"].iter_rows(min_row=5):
            if row[0].value == "FVMD":
                for cell in row[8:14]:
                    cell.value = 1200
            if row[0].value == "PDAH":
                for cell in row[8:14]:
                    cell.value = 100
        book.save(source)
        book.close()
        summary, output, import_output = self.run_updates([], source, self.pdah_floor_config())
        for path in (output, import_output):
            for column in range(9, 15):
                for key, rate in self.rates(path, column).items():
                    if key[0] == "PDAH":
                        expected = 170 if column == 14 and key[2] not in protected else 100
                        if "2026-10-30" <= key[2] <= "2026-11-01":
                            expected = max(expected, 110)
                        self.assertEqual(rate, expected)
        self.assertEqual(summary["mandatory_group_floor_change_count"], 26)

    def test_pdah_long_floor_dry_run_matches_export_and_detects_final_breach(self):
        source = self.workbook(airport_rate=100, city_rate=100)
        recommendations = self.root / "empty.json"
        recommendations.write_text('{"decisions": []}', encoding="utf-8")
        config = self.pdah_floor_config()
        source_bytes = source.read_bytes()
        preview = apply_updates(source, recommendations, None, config, None, True)
        written, output, _ = self.run_updates([], source, config)
        self.assertEqual(preview["mandatory_group_floor_change_count"], 2)
        self.assertEqual(preview["changes"], written["changes"])
        self.assertEqual(source.read_bytes(), source_bytes)
        book = openpyxl.load_workbook(output)
        sheet = book["Sheet1"]
        pdah_row = next(row[0].row for row in sheet.iter_rows(min_row=5) if row[0].value == "PDAH")
        sheet.cell(pdah_row, 14).value = 169.99
        checks = build_validation_rows(sheet, config, get_duration_columns(sheet, config), [], [])
        check = next(row for row in checks if row[0] == "Stawki ponizej obowiazkowego minimum klasy")
        self.assertEqual(check[2], 1)
        book.close()

    def test_pdah_floor_missing_duration_column_and_invalid_config_fail_closed(self):
        source = self.workbook(airport_rate=500, city_rate=500)
        book = openpyxl.load_workbook(source)
        book["Sheet1"].delete_cols(14)
        book.save(source)
        book.close()
        with self.assertRaisesRegex(ValueError, "Missing Excel duration columns"):
            self.run_updates([], source, self.pdah_floor_config())
        original = load_fixture_pricing_rules()
        for floors in ([], {"pdah": [{"minDays": 21, "maxDays": 35, "minimumRatePlnDay": 170}]},
                       {"PDAH": []}, {"PDAH": [{"minDays": 21, "maxDays": 35, "minimumRatePlnDay": -1}]},
                       {"PDAH": [{"minDays": 21, "maxDays": 36, "minimumRatePlnDay": 170}]}):
            path = self.root / "invalid-group-floors.json"
            path.write_text(json.dumps({"pricing": {**original["pricing"], "mandatoryGroupFloorsPlnDay": floors}}), encoding="utf-8")
            config = {**self.pdah_floor_config(), "pricing_rules_file": str(path)}
            with self.subTest(floors=floors), self.assertRaisesRegex(ValueError, "mandatoryGroupFloorsPlnDay"):
                self.run_updates([], self.workbook(airport_rate=500, city_rate=500), config)

    def test_pdah_floor_respects_higher_zone_minimum_without_double_counting(self):
        rules = load_fixture_pricing_rules()
        rules["pricing"]["mandatoryZoneFloorsPlnDay"]["SZLO"][-1]["minimumRatePlnDay"] = 200
        path = self.root / "higher-zone-floor.json"
        path.write_text(json.dumps(rules), encoding="utf-8")
        source = self.szlo_workbook(rate=500)
        book = openpyxl.load_workbook(source)
        for row in book["Sheet1"].iter_rows(min_row=5):
            if row[0].value == "PDAH":
                row[13].value = 100
        book.save(source)
        book.close()
        config = {**self.pdah_floor_config(), "pricing_rules_file": str(path)}
        summary, output, import_output = self.run_updates([], source, config)
        for output_path in (output, import_output):
            prices = self.rates(output_path, 14)
            self.assertEqual(prices["PDAH", "SZLO", "2026-10-01"], 200)
            self.assertEqual(prices["PDAH", "SZO1", "2026-10-01"], 170)
        self.assertEqual(summary["change_count"], 2)
        self.assertEqual(summary["mandatory_zone_floor_change_count"], 1)
        self.assertEqual(summary["mandatory_group_floor_change_count"], 1)

    def test_user_approved_baseline_allows_recommendations_but_not_calibration(self):
        manifest = {
            "status": "user_approved_baseline", "workbook_sha256": "b" * 64,
            "approved_by": "user", "approved_at": "2026-10-07",
        }
        path = self.root / "baseline.json"
        config = {**self.config, "baseline_manifest_file": str(path)}
        path.write_text(json.dumps(manifest), encoding="utf-8")
        result = load_baseline_confirmation(config, "b" * 64)
        self.assertTrue(result["recommendation_eligible"])
        self.assertFalse(result["confirmed"])
        self.assertFalse(result["calibration_eligible"])
        self.assertEqual(result["approved_by"], "user")
        for field in ("approved_by", "approved_at"):
            path.write_text(json.dumps({k: v for k, v in manifest.items() if k != field}), encoding="utf-8")
            with self.subTest(missing=field), self.assertRaises(ValueError):
                load_baseline_confirmation(config, "b" * 64)
        path.write_text(json.dumps(manifest), encoding="utf-8")
        with self.assertRaises(ValueError):
            load_baseline_confirmation(config, "c" * 64)

    def test_row_limit_is_optional_and_explicit_limits_are_still_enforced(self):
        book = openpyxl.Workbook()
        self.addCleanup(book.close)
        sheet = book.active
        sheet.cell(30001, 1, "CDMV")
        for config in (self.config, {"max_import_rows": None}, {}):
            with self.subTest(config=config):
                self.assertIsNone(get_import_row_limit(config))
                self.assertIsNone(validate_import_row_limit(sheet, config))
        self.assertEqual(validate_import_row_limit(sheet, {"max_import_rows": 31000}), 31000)
        with self.assertRaisesRegex(ValueError, "row limit exceeded"):
            validate_import_row_limit(sheet, {"max_import_rows": 28000})
        with self.assertRaises(ValueError):
            get_import_row_limit({"max_import_rows": 0})

    def test_approved_floor_corrected_baseline_is_not_claimed_as_imported(self):
        floors = [
            {"minDays": low, "maxDays": high, "minimumRatePlnDay": rate}
            for low, high, rate in ((1, 1, 300), (2, 2, 150), (3, 4, 130),
                                    (5, 7, 110), (8, 20, 90), (21, 35, 90))
        ]
        manifest = {
            "status": "user_approved_floor_correction", "workbook_sha256": "b" * 64,
            "approved_by": "user", "approved_at": "2026-10-05",
            "source_baseline": {"status": "confirmed_imported", "workbook_sha256": "a" * 64},
            "floor_corrections": {"SZLO": floors},
        }
        manifest_path = self.root / "baseline.json"
        config = {**self.config, "baseline_manifest_file": "baseline.json", "_config_dir": str(self.root),
                  "zone_location_labels": {"SZLO": "Szczecin Goleniow Airport (SZZ)"}}
        manifest_path.write_text(json.dumps(manifest), encoding="utf-8")
        result = load_baseline_confirmation(config, "b" * 64)
        self.assertTrue(result["recommendation_eligible"])
        self.assertFalse(result["confirmed"])
        self.assertFalse(result["calibration_eligible"])
        self.assertEqual(result["source_baseline"], manifest["source_baseline"])
        self.assertEqual(result["floor_corrections"], {"SZLO": floors})
        for active_floors in ({}, {"SZLO": [{**b, "minimumRatePlnDay": b["minimumRatePlnDay"] + 10} for b in floors]},
                              {"SZLO": [{**b, "minimumRatePlnDay": b["minimumRatePlnDay"] - 10} for b in floors]}):
            rules_path = self.root / "changed-rules.json"
            rules_path.write_text(json.dumps({"pricing": {"mandatoryZoneFloorsPlnDay": active_floors}}), encoding="utf-8")
            changed = load_baseline_confirmation({**config, "pricing_rules_file": str(rules_path)}, "b" * 64)
            self.assertEqual(changed["floor_corrections"], {"SZLO": floors})
            self.assertFalse(changed["calibration_eligible"])
        for field in ("approved_by", "approved_at", "source_baseline", "floor_corrections"):
            manifest_path.write_text(json.dumps({k: v for k, v in manifest.items() if k != field}), encoding="utf-8")
            with self.subTest(missing=field), self.assertRaises(ValueError):
                load_baseline_confirmation(config, "b" * 64)
        for correction in ({"OTHER": floors}, {"SZLO": floors[:-1]}, {}, {"SZLO": []}):
            manifest_path.write_text(json.dumps({**manifest, "floor_corrections": correction}), encoding="utf-8")
            with self.subTest(correction=correction), self.assertRaises(ValueError):
                load_baseline_confirmation(config, "b" * 64)
        manifest_path.write_text(json.dumps(manifest), encoding="utf-8")
        with self.assertRaises(ValueError):
            load_baseline_confirmation(config, "c" * 64)

    def test_explicit_rate_corrected_baseline_keeps_user_supplied_provenance(self):
        correction = {"group": "FVMD", "start_date": "2026-10-06", "end_date": "2026-10-12",
                      "rate_pln_day": 400, "all_duration_bands": True}
        manifest = {
            "status": "user_approved_rate_correction", "workbook_sha256": "b" * 64,
            "approved_by": "user", "approved_at": "2026-10-06",
            "source_baseline": {"status": "user_supplied", "workbook_sha256": "a" * 64},
            "rate_corrections": [correction],
        }
        path = self.root / "baseline.json"
        config = {**self.config, "baseline_manifest_file": "baseline.json", "_config_dir": str(self.root)}
        path.write_text(json.dumps(manifest), encoding="utf-8")
        result = load_baseline_confirmation(config, "b" * 64)
        self.assertTrue(result["recommendation_eligible"])
        self.assertFalse(result["confirmed"])
        self.assertFalse(result["calibration_eligible"])
        self.assertEqual(result["rate_corrections"], [correction])
        for field in ("approved_by", "approved_at", "source_baseline", "rate_corrections"):
            path.write_text(json.dumps({k: v for k, v in manifest.items() if k != field}), encoding="utf-8")
            with self.subTest(missing=field), self.assertRaises(ValueError):
                load_baseline_confirmation(config, "b" * 64)
        for invalid in ({"start_date": "bad"}, {"end_date": "2026-10-05"},
                        {"group": ""}, {"rate_pln_day": 0}, {"rate_pln_day": float("inf")},
                        {"all_duration_bands": False}):
            path.write_text(json.dumps({**manifest, "rate_corrections": [{**correction, **invalid}]}), encoding="utf-8")
            with self.subTest(invalid=invalid), self.assertRaises(ValueError):
                load_baseline_confirmation(config, "b" * 64)
        path.write_text(json.dumps(manifest), encoding="utf-8")
        with self.assertRaises(ValueError):
            load_baseline_confirmation(config, "c" * 64)

    def workbook(self, pickups=("2026-10-01",), airport_rate=40, city_rate=60):
        path = self.root / "baseline.xlsx"
        book = openpyxl.Workbook()
        sheet = book.active
        sheet.title = "Sheet1"
        sheet.append(["Rental rates for packages: INCLUSIVE POA"])
        sheet.append([None] * 8 + [1, 2, 3, 5, 8, 21])
        sheet.append([None] * 8 + [1, 2, 4, 7, 20, 35])
        sheet.append(["Group", "Description", "Rate code", "Zone", "Booking start date",
                      "Booking end date", "Pickup start date", "Pickup end date"] + ["Per day"] * 6)
        for pickup in pickups:
            for zone, rate in (("WA1", city_rate), ("WALO", airport_rate)):
                for group in ALLOWED + FROZEN:
                    adjustment = 1 if group in {"EDAV", "EDMV"} else 0
                    sheet.append([group, None, None, zone, pickup, pickup, pickup, pickup]
                                 + [rate + adjustment] * 6)
        sheet.freeze_panes = "I5"
        sheet["J5"].fill = PatternFill(fill_type="solid", fgColor="FFF2CC")
        sheet["J5"].comment = Comment("Baseline note", "Baseline")
        inherited = book.create_sheet("Inherited Notes")
        inherited["A1"] = "Keep inherited content"
        inherited.freeze_panes = "B3"
        book.save(path)
        book.close()
        return path

    def decision(self, days=2, pickup="2026-10-01", suggested=70, priority=False, **extra):
        return {
            "action": "increase",
            "recommendation_type": "force_top1_undercut" if priority else "top1_gap",
            "priority_rule_id": ("all-locations-autumn-2026-short" if days <= 4
                                 else "all-locations-autumn-2026-week") if priority else None,
            "location": "City", "start_date": pickup, "rental_days": days,
            "suggested_rate_pln_day": suggested, "maximum_import_rate_pln_day": suggested,
            "site_cap_rate_pln_day": suggested, "benchmark_rate_pln_day": suggested + 1,
            "mm_rate_pln_day": 60, "data_quality_status": "ok", "target_rank": 1,
            "scenario_id": f"{pickup}-{days}", **extra,
        }

    def run_updates(self, decisions, source=None, config=None):
        recommendations = self.root / "recommendations.json"
        recommendations.write_text(json.dumps({"decisions": decisions}), encoding="utf-8")
        output = self.root / "review.xlsx"
        import_output = self.root / "import.xlsx"
        summary = apply_updates(
            source or self.workbook(), recommendations, output, config or self.config,
            cli_groups=None, dry_run=False, import_output_path=import_output,
        )
        return summary, output, import_output

    def rates(self, path, column):
        book = openpyxl.load_workbook(path)
        try:
            return {(row[0].value, row[3].value, row[6].value): row[column - 1].value
                    for row in book["Sheet1"].iter_rows(min_row=5)}
        finally:
            book.close()

    def szlo_workbook(self, pickups=("2026-10-01",), rate=45):
        path = self.workbook(pickups=pickups, airport_rate=rate, city_rate=70)
        book = openpyxl.load_workbook(path)
        for row in book["Sheet1"].iter_rows(min_row=5):
            row[3].value = "SZLO" if row[3].value == "WALO" else "SZO1"
        book.save(path)
        book.close()
        return path

    def test_fvmd_is_omitted_only_in_szlo_for_all_dates_and_both_outputs(self):
        pickups = ("2026-10-01", "2026-10-31", "2027-01-10", "2027-01-11")
        source = self.szlo_workbook(pickups=pickups, rate=500)
        book = openpyxl.load_workbook(source)
        sheet = book["Sheet1"]
        sheet.append([" fvmd ", None, None, " szlo "] + ["2026-10-02"] * 4 + [500] * 6)
        sheet.row_dimensions[sheet.max_row].hidden = True
        for row in sheet.iter_rows(min_row=5):
            if row[0].value == "SWAV" and row[3].value == "SZLO":
                sheet.row_dimensions[row[0].row].height = 31
                sheet.row_dimensions[row[0].row].hidden = True
        book.save(source)
        book.close()
        before = source.read_bytes()
        for enabled in (False, True):
            config = {**self.config, "pickup_date_expansion": {
                "enabled": enabled, "start_date": "2026-10-01", "end_date": "2027-01-11",
            }}
            summary, output, import_output = self.run_updates([], source, config)
            for path in (output, import_output):
                prices = {(g, z, str(parse_date_value(d))): rate
                          for (g, z, d), rate in self.rates(path, 10).items()}
                self.assertFalse(any(g.strip().upper() == "FVMD" and z.strip().upper() == "SZLO"
                                     for g, z, _date in prices))
                self.assertEqual({d for g, z, d in prices if g == "FVMD" and z == "SZO1"}, set(pickups))
                for pickup in pickups:
                    self.assertEqual(prices["CFAV", "SZLO", pickup], 500)
                    self.assertEqual(prices["FVMD", "SZO1", pickup], 70)
                book = openpyxl.load_workbook(path)
                for row in book["Sheet1"].iter_rows(min_row=5):
                    if row[0].value == "SWAV" and row[3].value == "SZLO":
                        dimension = book["Sheet1"].row_dimensions[row[0].row]
                        self.assertEqual(dimension.height, 31)
                        self.assertTrue(dimension.hidden)
                book.close()
            self.assertEqual(summary["omitted_group_zone_rows"]["removed_row_count"], 5)
            self.assertEqual(source.read_bytes(), before)

    def test_zone_seeding_cannot_reintroduce_fvmd_in_szlo(self):
        source = self.workbook(airport_rate=500)
        book = openpyxl.load_workbook(source)
        for row in book["Sheet1"].iter_rows(min_row=5):
            if row[3].value == "WALO":
                row[3].value = "LOLO"
        book.save(source)
        book.close()
        config = {**self.config, "zone_seeds": {"SZLO": "LOLO"}}
        summary, output, import_output = self.run_updates([], source, config)
        for path in (output, import_output):
            prices = self.rates(path, 10)
            self.assertNotIn(("FVMD", "SZLO", "2026-10-01"), prices)
            self.assertEqual(prices["FVMD", "LOLO", "2026-10-01"], 500)
            self.assertEqual(prices["CFAV", "SZLO", "2026-10-01"], 500)
        self.assertEqual(summary["omitted_group_zone_rows"]["removed_row_count"], 1)

    def test_szlo_minima_cover_every_class_and_band_outside_protected_dates(self):
        pickups = ("2026-10-01", "2026-12-14", "2028-02-10")
        source = self.szlo_workbook(pickups=pickups)
        before = {column: self.rates(source, column) for column in range(9, 15)}
        summary, output, import_output = self.run_updates([], source)
        for path in (output, import_output):
            for column, minimum in ((9, 300), (10, 150), (11, 130), (12, 110), (13, 90), (14, 90)):
                prices = self.rates(path, column)
                for key, old_rate in before[column].items():
                    if key[0] == "FVMD" and key[1] == "SZLO":
                        self.assertNotIn(key, prices)
                        continue
                    self.assertEqual(prices[key], minimum if key[1] == "SZLO" else old_rate)
        self.assertEqual(summary["mandatory_zone_floor_change_count"], 198)
        checks = {row["check"]: row["issue_count"] for row in summary["validation"]}
        for check in ("Zmienione grupy wykluczone", "Zmienione klasy spoza listy dopuszczonej",
                      "Zmienione rekomendacje bez ceny benchmarku", "Stawki ponizej obowiazkowego minimum strefowego"):
            self.assertEqual(checks[check], 0)
        self.assertEqual(summary["recommendation_date_scope"]["date_count"], 0)
        self.assertEqual(summary["mandatory_zone_floor_date_scope"], {
            "start_date": "2026-10-01", "end_date": "2028-02-10", "date_count": 3,
        })
        book = openpyxl.load_workbook(output)
        changed_rows = [row for row in book["Changed Positions"].iter_rows(values_only=False)
                        if row[3].value == "SZLO"]
        self.assertTrue(changed_rows)
        self.assertTrue(all(row[0].fill.fgColor.rgb[-6:] == "FCE4D6" for row in changed_rows))
        for row in book["Recommendations Review"].iter_rows(min_row=2, values_only=True):
            self.assertEqual(row[1], "Gotowe z uwaga")
            self.assertIn("bez celu rankingowego", row[2])
            self.assertNotIn("top1", str(row[15]).lower())
            self.assertEqual(row[14], "Obowiazkowe minimum strefowe")
        book.close()
        for column in range(9, 15):
            self.assertEqual(self.rates(source, column), before[column])

    def test_protected_dates_preserve_all_szlo_rates_even_below_mandatory_floors(self):
        protected = {"2026-10-31", "2026-11-01", "2026-12-15", "2026-12-20", "2027-01-10"}
        unprotected = {"2026-10-30", "2026-11-02", "2026-12-14", "2027-01-11"}
        source = self.szlo_workbook(pickups=tuple(sorted(protected | unprotected)))
        source_bytes = source.read_bytes()
        before = {column: self.rates(source, column) for column in range(9, 15)}
        summary, output, import_output = self.run_updates([], source)
        self.assertEqual(summary["mandatory_zone_floor_change_count"], len(unprotected) * 11 * 6)
        self.assertTrue(all(row["status"] != "FAIL" for row in summary["validation"]))
        for path in (output, import_output):
            for column, minimum in ((9, 300), (10, 150), (11, 130), (12, 110), (13, 90), (14, 90)):
                for key, actual in self.rates(path, column).items():
                    expected = minimum if key[1] == "SZLO" and key[2] in unprotected else before[column][key]
                    self.assertEqual(actual, expected, f"{path.name}/{key}/{column}")
        self.assertEqual(source.read_bytes(), source_bytes)

    def test_validation_rejects_any_change_on_a_protected_date(self):
        source = self.szlo_workbook(pickups=("2026-12-20",))
        book = openpyxl.load_workbook(source)
        self.addCleanup(book.close)
        sheet = book["Sheet1"]
        change = {
            "group": "PDAH", "zone": "SZLO", "pickup_date": "2026-12-20",
            "duration_min_days": 2, "duration_max_days": 2, "duration_band": "2-2",
            "old_rate": 45, "new_rate": 150, "action": "increase",
            "recommendation_type": "mandatory_zone_floor",
        }
        checks = {row[0]: row for row in build_validation_rows(
            sheet, self.config, get_duration_columns(sheet, self.config), [change], [],
        )}
        self.assertEqual(checks["Zmienione stawki w chronionych datach"][1:3], ["FAIL", 1])
        self.assertEqual(checks["Zmienione grupy wykluczone"][2], 1)
        self.assertEqual(checks["Stawki ponizej obowiazkowego minimum strefowego"][2], 0)

    def test_szlo_higher_rates_and_other_szczecin_locations_remain_unchanged(self):
        source = self.szlo_workbook(rate=500)
        book = openpyxl.load_workbook(source)
        row = ["PDAH", None, None, "SZ1", "2026-10-01", "2026-10-01", "2026-10-01", "2026-10-01"]
        book["Sheet1"].append(row + [10] * 6)
        book.save(source)
        book.close()
        before = {column: self.rates(source, column) for column in range(9, 15)}
        summary, output, import_output = self.run_updates([], source)
        self.assertEqual(summary["change_count"], 0)
        for path in (output, import_output):
            for column in range(9, 15):
                expected = {key: value for key, value in before[column].items()
                            if key[0] != "FVMD" or key[1] != "SZLO"}
                self.assertEqual(self.rates(path, column), expected)

    def test_szlo_mandatory_minimum_overrides_priority_top1_and_long_duration_floor(self):
        source = self.szlo_workbook()
        config = {**self.config, "location_zones": {"Airport": ["SZLO"]}}
        decisions = [self.decision(day, location="Airport", suggested=30, priority=day <= 7)
                     for day in (2, 3, 4, 5, 6, 7, *range(8, 15))]
        _, output, _ = self.run_updates(decisions, source, config)
        for column, minimum in ((10, 150), (11, 130), (12, 110), (13, 90)):
            rates = self.rates(output, column)
            self.assertEqual(rates["CDMV", "SZLO", "2026-10-01"], minimum)
            self.assertGreaterEqual(rates["EDMV", "SZLO", "2026-10-01"], minimum)

    def test_szlo_dry_run_calculates_floors_without_writing_files_or_changing_source(self):
        source = self.szlo_workbook()
        recommendations = self.root / "empty.json"
        recommendations.write_text('{"decisions": []}', encoding="utf-8")
        before = source.read_bytes()
        summary = apply_updates(source, recommendations, self.root / "unused.xlsx", self.config,
                                cli_groups=None, dry_run=True, import_output_path=self.root / "unused-import.xlsx")
        self.assertEqual(summary["mandatory_zone_floor_change_count"], 66)
        self.assertEqual(source.read_bytes(), before)
        self.assertFalse((self.root / "unused.xlsx").exists())
        self.assertFalse((self.root / "unused-import.xlsx").exists())

    def test_invalid_mandatory_floor_configuration_cannot_silently_export_rates(self):
        original = load_fixture_pricing_rules()
        bands = original["pricing"]["mandatoryZoneFloorsPlnDay"]["SZLO"]
        invalid_tables = [None, [], {"szlo": bands}, {"SZLO": bands[:-1]},
                          {"SZLO": [{**bands[0], "minimumRatePlnDay": 0}, *bands[1:]]},
                          {"SZLO": [{**bands[0], "minimumRatePlnDay": "300"}, *bands[1:]]},
                          {"SZLO": [{**bands[0], "maxDays": 2}, *bands[1:]]}]
        for invalid in invalid_tables:
            with self.subTest(invalid=invalid):
                rules = {"pricing": {**original["pricing"], "mandatoryZoneFloorsPlnDay": invalid}}
                path = self.root / "bad-floors.json"
                path.write_text(json.dumps(rules), encoding="utf-8")
                with self.assertRaisesRegex(ValueError, "mandatoryZoneFloorsPlnDay"):
                    self.run_updates([], self.szlo_workbook(), {**self.config, "pricing_rules_file": str(path)})

    def test_validation_detects_a_remaining_szlo_floor_breach_and_rejects_false_exceptions(self):
        source = self.szlo_workbook(rate=500)
        book = openpyxl.load_workbook(source)
        sheet = book["Sheet1"]
        row = next(row for row in sheet.iter_rows(min_row=5)
                   if row[0].value == "PDAH" and row[3].value == "SZLO")
        row[9].value = 149
        fake_changes = [{"group": "PDAH", "zone": zone, "pickup_date": "2026-10-01",
                         "duration_min_days": 2, "duration_max_days": 2, "duration_band": "2",
                         "recommendation_type": "mandatory_zone_floor", "action": "increase",
                         "old_rate": 100, "new_rate": new_rate, "cell": cell}
                        for zone, new_rate, cell in (("SZO1", 150, "J5"), ("SZLO", 149, row[9].coordinate))]
        validation = build_validation_rows(sheet, self.config, get_duration_columns(sheet, self.config), fake_changes, [])
        checks = {row[0]: row[2] for row in validation}
        self.assertEqual(checks["Stawki ponizej obowiazkowego minimum strefowego"], 1)
        self.assertEqual(checks["Zmienione grupy wykluczone"], 2)
        self.assertEqual(checks["Zmienione klasy spoza listy dopuszczonej"], 2)
        book.close()

    def test_validation_retains_existing_partial_parity_record_contract(self):
        book = openpyxl.load_workbook(self.workbook())
        sheet = book["Sheet1"]
        rows = build_validation_rows(sheet, self.config, get_duration_columns(sheet, self.config),
                                     [{"recommendation_type": "group_parity", "cell": "J5"}], [])
        benchmark_check = next(row for row in rows if row[0] == "Zmienione rekomendacje bez ceny benchmarku")
        self.assertEqual(benchmark_check[2], 0)
        book.close()

    def test_missing_duration_header_cannot_skip_a_mandatory_szlo_floor(self):
        source = self.szlo_workbook()
        book = openpyxl.load_workbook(source)
        book["Sheet1"]["J2"] = None
        book.save(source)
        book.close()
        with self.assertRaisesRegex(ValueError, "Missing Excel duration.*2"):
            self.run_updates([], source)
        self.assertFalse((self.root / "import.xlsx").exists())

    def test_szlo_active_market_dry_run_matches_written_changes(self):
        source = self.szlo_workbook()
        book = openpyxl.load_workbook(source)
        for row in book["Sheet1"].iter_rows(min_row=5):
            if row[3].value == "SZLO" and row[0].value in {"EDAV", "EDMV"}:
                row[9].value = 45
        book.save(source)
        book.close()
        config = {**self.config, "location_zones": {"Airport": ["SZLO"]}}
        written, _, _ = self.run_updates([self.decision(location="Airport", suggested=30, priority=True)], source, config)
        before = source.read_bytes()
        preview = apply_updates(source, self.root / "recommendations.json", None, config, None, True)
        self.assertEqual(source.read_bytes(), before)
        self.assertEqual(preview["change_statistics"], written["change_statistics"])
        self.assertEqual(preview["mandatory_zone_floor_change_count"], written["mandatory_zone_floor_change_count"])
        self.assertEqual(preview["changes"], written["changes"])

    def test_final_floor_validation_reads_actual_cells_not_only_planned_prices(self):
        book = openpyxl.load_workbook(self.szlo_workbook(rate=500))
        sheet = book["Sheet1"]
        row = next(row for row in sheet.iter_rows(min_row=5) if row[3].value == "SZLO")
        row[9].value = 149
        planned = [{"cell": row[9].coordinate, "new_rate": 150}]
        checks = build_validation_rows(sheet, self.config, get_duration_columns(sheet, self.config), planned, [])
        floor_check = next(check for check in checks if check[0] == "Stawki ponizej obowiazkowego minimum strefowego")
        self.assertEqual(floor_check[2], 1)
        book.close()

    def test_priority_wins_city_cap_and_normal_floor_for_all_priority_bands(self):
        for days, column, expected in (([2], 10, 30), ([3, 4], 11, 30), ([5, 6, 7], 12, 40)):
            with self.subTest(days=days):
                source = self.workbook(airport_rate=10, city_rate=35)
                decisions = [self.decision(day, suggested=29, priority=True, action="decrease") for day in days]
                summary, output, _ = self.run_updates(decisions, source)
                self.assertEqual(summary["city_top1_airport_cap_scope_count"], 0)
                self.assertEqual(summary["city_top1_airport_cap_conflict_count"], 0)
                prices = self.rates(output, column)
                for group in ALLOWED:
                    self.assertEqual(prices[group, "WA1", "2026-10-01"], expected + (group in {"EDAV", "EDMV"}))

    def test_normal_floor_conflict_ignores_only_affected_band_and_exports_market_target(self):
        config = merge_fixture_config({**self.config, "minimum_rates": {
            "global_min_pln_day": 50, "bands": [],
            "scoped_overrides": [{"zones": ["WA1"], "groups": ALLOWED,
                                  "min_days": 2, "max_days": 2, "min_pln_day": 53}],
        }})
        source = self.workbook()
        original = source.read_bytes()
        summary, output, import_output = self.run_updates([self.decision(day) for day in (2, 3, 4)], source, config)
        self.assertEqual(summary["city_top1_airport_cap_conflict_count"], 0)
        self.assertEqual(summary["city_top1_airport_cap_conflicts"], [])
        self.assertEqual(summary["skipped_target_count"], 0)
        self.assertEqual(summary["city_top1_airport_cap_ignored_count"], 1)
        scope, = summary["city_top1_airport_cap_ignored_scopes"]
        self.assertEqual((scope["zone"], scope["pickup_date"], scope["duration_band"]), ("WA1", "2026-10-01", "2"))
        self.assertEqual(scope["groups"], sorted(ALLOWED))
        self.assertIn("53", scope["reason"])
        self.assertIn("52", scope["reason"])
        self.assertEqual(summary["city_top1_airport_cap_violation_count"], 0)
        for change in summary["changes"]:
            ignored = change["duration_band"] == "2"
            self.assertEqual(change["city_top1_airport_cap_ignored"], ignored)
            self.assertEqual(change["city_top1_airport_cap_active"], not ignored)
            self.assertEqual(change["city_top1_airport_cap_applied"], not ignored)
            self.assertEqual(change["target_rank"], 1)
            if ignored:
                self.assertEqual(change["city_top1_airport_cap_ignored_reason"], scope["reason"])
        for path in (output, import_output):
            prices = self.rates(path, 10)
            remaining = self.rates(path, 11)
            for group in ALLOWED:
                adjustment = group in {"EDAV", "EDMV"}
                self.assertEqual(prices[group, "WA1", "2026-10-01"], 70 + adjustment)
                self.assertEqual(remaining[group, "WA1", "2026-10-01"], 52 + adjustment)
            book = openpyxl.load_workbook(path)
            self.assertNotEqual(book["Sheet1"]["J5"].fill.fgColor.rgb[-6:], "FFF2CC")
            if path == import_output:
                self.assertIsNone(book["Sheet1"]["J5"].comment)
            else:
                self.assertIn("Nowa stawka: 70 PLN", book["Sheet1"]["J5"].comment.text)
                self.assertNotIn("Limit oddzialu miejskiego", book["Sheet1"]["J5"].comment.text)
                self.assertNotEqual(book["Sheet1"]["J5"].comment.text, "Baseline note")
            book.close()
        preview = apply_updates(source, self.root / "recommendations.json", None, config, None, True)
        for field in ("changes", "change_statistics", "skipped_target_count", "city_top1_airport_cap_ignored_count",
                      "city_top1_airport_cap_ignored_scopes", "city_top1_airport_cap_conflict_count", "city_top1_airport_cap_conflicts"):
            self.assertEqual(preview[field], summary[field], field)
        self.assertEqual(source.read_bytes(), original)

    def test_cap_ignored_count_counts_duration_bands_not_affected_classes(self):
        config = merge_fixture_config({**self.config, "minimum_rates": {
            "global_min_pln_day": 53, "bands": [], "scoped_overrides": [],
        }})
        source = self.workbook()
        original = source.read_bytes()
        summary, output, imported = self.run_updates([self.decision(day) for day in (2, 3, 4)], source, config)
        self.assertEqual(summary["city_top1_airport_cap_ignored_count"], 2)
        scopes = summary["city_top1_airport_cap_ignored_scopes"]
        self.assertEqual({scope["duration_band"] for scope in scopes}, {"2", "3-4"})
        self.assertEqual(len(scopes), 2)
        for scope in scopes:
            self.assertEqual(scope["groups"], sorted(ALLOWED))
            self.assertTrue(scope["reason"])
        self.assertEqual(summary["city_top1_airport_cap_conflict_count"], 0)
        self.assertEqual(summary["city_top1_airport_cap_conflicts"], [])
        self.assertEqual(summary["skipped_target_count"], 0)
        for path in (output, imported):
            for col in (10, 11):
                prices = self.rates(path, col)
                for group in ALLOWED:
                    self.assertEqual(prices[group, "WA1", "2026-10-01"], 70 + (group in {"EDAV", "EDMV"}))
        self.assertEqual(source.read_bytes(), original)

    def test_priority_expires_after_october_25_and_normal_cap_resumes(self):
        source = self.workbook(pickups=("2026-10-25", "2026-10-26"))
        summary, output, _ = self.run_updates([
            self.decision(pickup="2026-10-25", suggested=80, priority=True),
            self.decision(pickup="2026-10-26", suggested=80),
        ], source)
        prices = self.rates(output, 10)
        self.assertEqual(prices["CDMV", "WA1", "2026-10-25"], 79)
        self.assertEqual(prices["CDMV", "WA1", "2026-10-26"], 52)
        self.assertEqual(summary["city_top1_airport_cap_scope_count"], 1)

    def test_stale_priority_id_does_not_disable_normal_cap(self):
        source = self.workbook(pickups=("2026-10-26",), airport_rate=60)
        summary, output, _ = self.run_updates([self.decision(pickup="2026-10-26", priority=True)], source)
        self.assertEqual(summary["city_top1_airport_cap_scope_count"], 1)
        self.assertEqual(self.rates(output, 10)["CDMV", "WA1", "2026-10-26"], 60)

    def test_holidays_protect_all_rates_even_with_priority_tag(self):
        pickups = ("2026-10-31", "2026-11-01", "2026-12-15", "2027-01-10")
        source = self.workbook(pickups=pickups)
        before = {column: self.rates(source, column) for column in range(9, 15)}
        summary, output, import_output = self.run_updates([
            self.decision(pickup=pickup, priority=True, suggested=20, action="decrease") for pickup in pickups
        ], source)
        self.assertEqual(summary["change_count"], 0)
        for path in (output, import_output):
            for column in range(9, 15):
                self.assertEqual(self.rates(path, column), before[column])

    def test_premium_plus_one_and_excluded_classes_stay_frozen(self):
        source = self.workbook(city_rate=45)
        before = {column: self.rates(source, column) for column in range(9, 15)}
        _, output, import_output = self.run_updates([self.decision(suggested=50, priority=True)], source)
        for path in (output, import_output):
            prices = self.rates(path, 10)
            for group in ALLOWED:
                self.assertEqual(prices[group, "WA1", "2026-10-01"], 49 + (group in {"EDAV", "EDMV"}))
            for column in range(9, 15):
                prices = self.rates(path, column)
                for key, rate in before[column].items():
                    if key[0] in FROZEN or key[1] == "WALO" or column != 10:
                        self.assertEqual(prices[key], rate)

    def test_approved_8_to_14_evidence_is_not_missing_required_15_to_20(self):
        summary, output, _ = self.run_updates([
            self.decision(day, suggested=45, site_cap_rate_pln_day=50, action="decrease") for day in range(8, 15)
        ])
        incomplete = next(row for row in summary["validation"] if row["check"] == "Niepelne pokrycie przedzialu duration")
        self.assertEqual(incomplete["issue_count"], 0)
        self.assertEqual(summary["skipped_target_count"], 0)
        self.assertEqual(self.rates(output, 13)["CDMV", "WA1", "2026-10-01"], 45)
        book = openpyxl.load_workbook(output)
        for row in book["Recommendations Review"].iter_rows(min_row=2, values_only=True):
            self.assertNotEqual(row[1], "Sprawdz")
            self.assertNotIn("brak danych dla duration", row[2])
        book.close()

    def test_missing_required_14_day_evidence_blocks_whole_8_to_20_band(self):
        summary, output, _ = self.run_updates([self.decision(day, suggested=45, action="decrease") for day in range(8, 14)])
        self.assertEqual(summary["change_count"], 0)
        skipped = next(row for row in summary["validation"] if row["check"] == "Pominiete rekomendacje")
        self.assertIn("lacks scenarios for 14", skipped["details"])
        self.assertEqual(self.rates(output, 13)["EDAV", "WA1", "2026-10-01"], 61)

    def test_invalid_evidence_is_not_an_expected_floor_rank_limit(self):
        decisions = [self.decision(day, suggested=45, action="decrease") for day in range(8, 15)]
        decisions[0]["data_quality_status"] = "missing_mm_rate"
        summary, output, _ = self.run_updates(decisions)
        self.assertEqual(summary["change_count"], 0)
        skipped = next(row for row in summary["validation"] if row["check"] == "Pominiete rekomendacje")
        self.assertIn("missing or invalid constraint data", skipped["details"])
        self.assertEqual(self.rates(output, 13)["CDMV", "WA1", "2026-10-01"], 60)

    def test_expected_floor_and_premium_rank_limits_are_notes_not_bad_evidence(self):
        config = merge_fixture_config({**self.config, "city_top1_airport_cap": {"enabled": False}})
        summary, output, _ = self.run_updates([self.decision(suggested=45, action="decrease")], config=config)
        self.assertTrue(all(not change["target_achievable"] for change in summary["changes"]))
        book = openpyxl.load_workbook(output)
        rows = list(book["Recommendations Review"].iter_rows(min_row=2, values_only=True))
        self.assertTrue(rows)
        for row in rows:
            self.assertEqual(row[1], "Gotowe z uwaga")
            self.assertIn("oczekiwane ograniczenie", row[2])
        book.close()
        change = dict(summary["changes"][0], ranking_limit_reason="", benchmark_rate=None)
        self.assertEqual(get_review_status([change]), "Sprawdz")

    def test_premium_only_rank_limit_is_not_reported_as_missing_evidence(self):
        config = merge_fixture_config({**self.config, "minimum_rates": {"global_min_pln_day": 0, "bands": []},
                               "city_top1_airport_cap": {"enabled": False}})
        summary, _, _ = self.run_updates([self.decision(suggested=55, action="decrease")], config=config)
        base = next(change for change in summary["changes"] if change["group"] == "CDMV")
        premium = next(change for change in summary["changes"] if change["group"] == "EDAV")
        self.assertTrue(base["target_achievable"])
        self.assertFalse(premium["target_achievable"])
        self.assertEqual(premium["ranking_limit_reason"], "premium")
        self.assertEqual(get_review_status([premium]), "Gotowe z uwaga")

    def test_floor_aware_top2_fallback_is_explicit_in_review(self):
        source = self.workbook(city_rate=45)
        _, output, _ = self.run_updates([
            self.decision(suggested=50, priority=True, target_rank=2, recommendation_type="priority_top3")
        ], source)
        book = openpyxl.load_workbook(output)
        for row in book["Recommendations Review"].iter_rows(min_row=2, values_only=True):
            self.assertEqual(row[1], "Gotowe z uwaga")
            self.assertIn("top2", row[2])
            self.assertIn("oczekiwane ograniczenie", row[2])
        book.close()

    def test_generated_legend_exposes_active_hierarchy_priority_and_protection(self):
        _, output, _ = self.run_updates([])
        book = openpyxl.load_workbook(output)
        self.assertEqual(book["Changed Positions"]["A5"].value, "Hierarchia i scalanie")
        self.assertGreater(book["Changed Positions"].row_dimensions[5].height or 15, 100)
        legend = " ".join(str(cell.value or "") for row in book["Changed Positions"] for cell in row)
        for required in ("Hierarchia", "nadrzedny", "2026-10-25", "2-2", "3-4", "5-7",
                         "30 PLN", "40 PLN", "2026-10-31", "2027-01-10", "EDAV=baza+1", *ALLOWED):
            self.assertIn(required, legend)
        book.close()

    def test_freeze_cleanup_keeps_unfrozen_selections_and_split_views(self):
        source = self.workbook()
        book = openpyxl.load_workbook(source)
        sheet = book.create_sheet("Unfrozen view")
        sheet.sheet_view.selection[0].activeCell = "D7"
        sheet.sheet_view.selection[0].sqref = "D7"
        split = book.create_sheet("Split view")
        split.sheet_view.pane = openpyxl.worksheet.views.Pane(state="split", xSplit=2000)
        book.save(source)
        book.close()
        _, output, _ = self.run_updates([], source)
        book = openpyxl.load_workbook(output)
        self.assertEqual(book["Unfrozen view"].sheet_view.selection[0].activeCell, "D7")
        self.assertEqual(book["Split view"].sheet_view.pane.state, "split")
        book.close()

    def test_config_cannot_reverse_approved_priority_over_city_cap(self):
        rules = load_fixture_pricing_rules()
        rules["pricing"]["pricingPolicy"] = {
            "precedence": ["protected_dates", "class_allowlist", "required_evidence", "city_airport_cap",
                           "priority_top1", "scoped_floor", "zone_floor", "seasonal_floor", "ranking_target", "premium_parity"],
            "normalFloorCapConflict": "ignore_cap",
        }
        path = self.root / "invalid-hierarchy.json"
        path.write_text(json.dumps(rules), encoding="utf-8")
        config = {**self.config, "pricing_rules_file": str(path)}
        with self.assertRaisesRegex(ValueError, "pricingPolicy"):
            self.run_updates([self.decision(priority=True)], config=config)

    def test_retired_baseline_preservation_and_unknown_conflict_modes_fail_closed(self):
        original = load_fixture_pricing_rules()
        source = self.workbook()
        source_bytes = source.read_bytes()
        for mode in ("preserve_baseline", "unknown", None):
            rules = json.loads(json.dumps(original))
            rules["pricing"]["pricingPolicy"]["normalFloorCapConflict"] = mode
            path = self.root / "invalid-conflict-policy.json"
            path.write_text(json.dumps(rules), encoding="utf-8")
            with self.subTest(mode=mode), self.assertRaisesRegex(ValueError, "pricingPolicy"):
                self.run_updates([self.decision()], source, {**self.config, "pricing_rules_file": str(path)})
            self.assertFalse((self.root / "review.xlsx").exists())
            self.assertFalse((self.root / "import.xlsx").exists())
            self.assertEqual(source.read_bytes(), source_bytes)

    def test_both_outputs_clear_inherited_and_generated_frozen_panes(self):
        source = self.workbook()
        _, output, import_output = self.run_updates([self.decision(priority=True)], source)
        for path in (output, import_output):
            book = openpyxl.load_workbook(path)
            for sheet in book:
                self.assertIsNone(sheet.freeze_panes, f"{path.name}/{sheet.title}")
                self.assertFalse(any(pane.state in {"frozen", "frozenSplit"}
                                     for view in sheet.views.sheetView for pane in [view.pane] if pane))
            if path == output:
                self.assertEqual(book["Inherited Notes"]["A1"].value, "Keep inherited content")
            else:
                self.assertEqual(book.sheetnames, ["Sheet1"])
            book.close()
        original = openpyxl.load_workbook(source)
        self.assertEqual(original["Sheet1"].freeze_panes, "I5")
        original.close()


class ImportPackageTests(unittest.TestCase):
    def test_package_validation_keeps_percent_escaped_part_names(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            book = openpyxl.Workbook()
            book.active["A1"].comment = Comment("Comment", "Test")
            original = root / "original.xlsx"
            book.save(original)
            with ZipFile(original) as archive:
                entries = {entry.filename: archive.read(entry) for entry in archive.infolist()}
            drawing = "xl/drawings/commentsDrawing1.vml"
            for filename in ("comment%20drawing.vml", "comment%23drawing.vml"):
                for absolute in (True, False):
                    for present in (True, False):
                        with self.subTest(filename=filename, absolute=absolute, present=present):
                            package = root / "escaped.xlsx"
                            renamed = "xl/drawings/" + filename
                            with ZipFile(package, "w") as archive:
                                for name, data in entries.items():
                                    if name == drawing:
                                        if not present:
                                            continue
                                        name = renamed
                                    elif name == "xl/worksheets/_rels/sheet1.xml.rels":
                                        relationships = ElementTree.fromstring(data)
                                        for relationship in relationships:
                                            if relationship.get("Type", "").endswith("/vmlDrawing"):
                                                relationship.set("Target", "/" + renamed if absolute else "../drawings/" + filename)
                                        data = ElementTree.tostring(relationships)
                                    archive.writestr(name, data)
                            if present:
                                excel_updater.validate_workbook_relationships(package)
                            else:
                                with self.assertRaisesRegex(ValueError, "missing part"):
                                    excel_updater.validate_workbook_relationships(package)

    def test_package_validation_rejects_absolute_and_relative_missing_parts(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            book = openpyxl.Workbook()
            book.active["A1"].comment = Comment("Comment", "Test")
            valid = root / "valid.xlsx"
            book.save(valid)
            excel_updater.validate_workbook_relationships(valid)
            with ZipFile(valid) as archive:
                entries = {entry.filename: archive.read(entry) for entry in archive.infolist()}
            drawing = "xl/drawings/commentsDrawing1.vml"
            self.assertIn(drawing, entries)
            for target in ("/xl/drawings/commentsDrawing1.vml", "../drawings/commentsDrawing1.vml"):
                with self.subTest(target=target):
                    broken = root / "broken.xlsx"
                    with ZipFile(broken, "w") as archive:
                        for name, data in entries.items():
                            if name == drawing:
                                continue
                            if name == "xl/worksheets/_rels/sheet1.xml.rels":
                                relationships = ElementTree.fromstring(data)
                                for relationship in relationships:
                                    if relationship.get("Type", "").endswith("/vmlDrawing"):
                                        relationship.set("Target", target)
                                data = ElementTree.tostring(relationships)
                            archive.writestr(name, data)
                    with self.assertRaisesRegex(ValueError, "missing part.*commentsDrawing1.vml"):
                        excel_updater.validate_workbook_relationships(broken)

    def test_clean_import_keeps_valid_chart_and_external_hyperlink(self):
        from openpyxl.chart import BarChart, Reference

        with tempfile.TemporaryDirectory() as directory:
            book = openpyxl.Workbook()
            book.active.title = "Sheet1"
            book.active.append([35, 60])
            book.active["A1"].comment = Comment("Rate", "Test")
            book.active["B1"].hyperlink = "https://www.discovercars.com/?a=1&b=2"
            chart = BarChart()
            chart.add_data(Reference(book.active, min_col=1, max_col=2, min_row=1))
            book.active.add_chart(chart, "D5")
            book.save(Path(directory) / "recommendations.xlsx")
            imported = Path(directory) / "import.xlsx"
            save_import_ready_workbook(book, "Sheet1", imported)
            excel_updater.validate_workbook_relationships(imported)
            result = openpyxl.load_workbook(imported)
            self.assertEqual(len(result.active._charts), 1)
            self.assertEqual(result.active["B1"].hyperlink.target, "https://www.discovercars.com/?a=1&b=2")
            self.assertEqual(list(result.active.values)[0], (35, 60))
            result.close()

    def test_clean_import_after_recommendation_save_has_no_missing_drawing(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            book = openpyxl.Workbook()
            sheet = book.active
            sheet.title = "Sheet1"
            sheet.merge_cells("A1:H1")
            sheet["A1"] = "Rental rates for packages: INCLUSIVE POA"
            sheet["J5"] = 60.25
            sheet["J5"].fill = PatternFill("solid", fgColor="00FF00")
            sheet["J5"].comment = Comment("Previous rate: 35 PLN", "Test")
            sheet["A5"] = "CDMV"
            sheet["A5"].hyperlink = "https://www.discovercars.com/"
            book.create_sheet("Changed Positions")["O2"].comment = Comment("Change", "Test")
            recommendations = root / "recommendations.xlsx"
            book.save(recommendations)
            original_recommendations = recommendations.read_bytes()
            imported = root / "import.xlsx"

            save_import_ready_workbook(book, "Sheet1", imported)

            with ZipFile(imported) as archive:
                names = set(archive.namelist())
                for name in names:
                    if not name.endswith(".rels"):
                        continue
                    for relationship in ElementTree.fromstring(archive.read(name)):
                        if relationship.get("TargetMode") == "External":
                            continue
                        target = relationship.get("Target", "")
                        resolved = (target.lstrip("/") if target.startswith("/") else
                                    posixpath.normpath(posixpath.join(posixpath.dirname(posixpath.dirname(name)), target)))
                        self.assertIn(resolved, names, f"{name}: dangling target {target}")
                self.assertFalse(any("comments" in name.lower() or name.endswith(".vml") for name in names))
                sheet_xml = ElementTree.fromstring(archive.read("xl/worksheets/sheet1.xml"))
                self.assertIsNone(sheet_xml.find("{http://schemas.openxmlformats.org/spreadsheetml/2006/main}legacyDrawing"))
            result = openpyxl.load_workbook(imported)
            self.assertEqual(result.sheetnames, ["Sheet1"])
            self.assertEqual(result.active["J5"].value, 60.25)
            self.assertEqual(result.active["J5"].fill.patternType, "solid")
            self.assertEqual(result.active["J5"].fill.fgColor.rgb, "0000FF00")
            self.assertEqual(str(result.active.merged_cells), "A1:H1")
            self.assertIsNone(result.active["J5"].comment)
            self.assertEqual(result.active["A5"].hyperlink.target, "https://www.discovercars.com/")
            result.close()
            self.assertEqual(recommendations.read_bytes(), original_recommendations)
            with ZipFile(recommendations) as archive:
                self.assertIn("xl/drawings/commentsDrawing1.vml", archive.namelist())


class LocationTypeFloorTests(unittest.TestCase):
    workbook = PricingPolicyTests.workbook
    decision = PricingPolicyTests.decision
    run_updates = PricingPolicyTests.run_updates
    rates = PricingPolicyTests.rates

    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.config = load_fixture_config()
        self.config.update(baseline_manifest_file="", pickup_date_expansion={"enabled": False}, zone_seeds={})
        self.config["pricing_rules_file"] = str(FIXTURE_PRICING_RULES_PATH)
        self.config["location_zones"]["City"] = ["WA1"]

    def test_fixture_location_type_rule_replaces_priority_scoped_and_seasonal_floors_at_endpoints(self):
        for pickup in ("2026-07-01", "2026-10-25", "2026-10-26", "2028-02-10"):
            for zone, minimum in (("WA1", 60), ("WR2", 60), ("WALO", 30), ("KRTI", 30)):
                for low, high in ((1, 1), (2, 2), (3, 4), (5, 7), (8, 20)):
                    for group in ALLOWED:
                        target = self.decision(low, pickup, priority=True)
                        target.update(zone=zone, group=group, target_date=parse_date_value(pickup),
                                      duration_min_days=low, duration_max_days=high)
                        with self.subTest(pickup=pickup, zone=zone, group=group, band=(low, high)):
                            self.assertEqual(get_minimum_rate(target, self.config)[0], minimum)
                            self.assertEqual(get_mandatory_floor(zone, group, low, high, self.config, pickup)[0], minimum)
        self.assertIsNone(get_mandatory_floor("WALO", "CDMV", 21, 35, self.config, "2026-10-26")[0])
        self.assertEqual(get_mandatory_floor("SZLO", "CDMV", 1, 1, self.config, "2026-10-26")[0], 300)

    def test_every_registry_zone_and_band_corrects_missing_market_without_lowering(self):
        source = self.workbook(pickups=(), city_rate=10)
        registry = json.loads(FIXTURE_REGISTRY_PATH.read_text(encoding="utf-8"))
        expected = {}
        book = openpyxl.load_workbook(source)
        book["Sheet1"].delete_rows(5, book["Sheet1"].max_row)
        for location in registry["locations"]:
            for zone in location["zones"]:
                for group in ALLOWED + FROZEN + ["ZZAV"]:
                    for pickup, old in (("2026-10-13", 10), ("2028-02-10", 1200)):
                        if group == "FVMD" and zone == "SZLO":
                            continue
                        book["Sheet1"].append([group, None, None, zone] + [pickup] * 4 + [old] * 6)
                        rates = []
                        for index, zone_min in enumerate((300, 150, 130, 110, 90, 90)):
                            minimum = zone_min if zone == "SZLO" else 0
                            if group == "PDAH" and index == 5:
                                minimum = max(minimum, 170)
                            if group in ALLOWED and index < 5 and zone != "SZLO":
                                minimum = (60 if location["kind"] == "city" else 30) + (group in {"EDAV", "EDMV"})
                            rates.append(max(old, minimum))
                        expected[group, zone, pickup] = rates
        book.save(source)
        book.close()
        original = source.read_bytes()
        summary, output, imported = self.run_updates([], source)
        for path in (output, imported):
            book = openpyxl.load_workbook(path)
            for row in book["Sheet1"].iter_rows(min_row=5, values_only=True):
                key = row[0], row[3], row[6]
                self.assertEqual(list(row[8:14]), expected[key], key)
            self.assertTrue(all(sheet.freeze_panes is None for sheet in book))
            if path == imported:
                self.assertEqual(book.sheetnames, ["Sheet1"])
                self.assertTrue(all(cell.comment is None for row in book["Sheet1"] for cell in row))
            else:
                legend = " ".join(str(c.value or "") for row in book["Changed Positions"] for c in row)
                for text in ("60 PLN", "30 PLN", "1-20", "typ", "zastepuje", "bez celu rankingowego"):
                    self.assertIn(text, legend)
            book.close()
        self.assertEqual(source.read_bytes(), original)
        self.assertEqual(summary["change_statistics"]["decrease_count"], 0)
        self.assertEqual(summary["recommendation_date_scope"]["date_count"], 0)
        self.assertEqual(summary["mandatory_zone_floor_date_scope"]["start_date"], "2026-10-13")
        self.assertTrue(all(change["target_rank"] == "" for change in summary["changes"]))
        self.assertFalse(any(row["status"] == "FAIL" for row in summary["validation"]))

    def test_floor_protects_real_rank_fail_hold_and_invalid_market_data(self):
        for extra in ({"action": "hold", "recommendation_type": "priority_top3_unachievable", "target_achievable": False},
                      {"action": "hold", "suggested_rate_pln_day": None, "maximum_import_rate_pln_day": None,
                       "data_quality_status": "missing_top3"},
                      {"data_quality_status": "missing_mm_rate"}):
            with self.subTest(extra=extra):
                source = self.workbook(city_rate=10, airport_rate=10)
                summary, output, _ = self.run_updates([self.decision(priority=True, **extra)], source)
                self.assertEqual(self.rates(output, 10)["CDMV", "WA1", "2026-10-01"], 60)
                self.assertEqual(self.rates(output, 10)["EDMV", "WA1", "2026-10-01"], 61)
                self.assertTrue(all(change["target_rank"] == "" for change in summary["changes"]))
                self.assertTrue(all(change["recommendation_type"] == "mandatory_zone_floor"
                                    for change in summary["changes"] if change["group"] in ALLOWED))

    def test_new_floor_respects_protected_dates_and_existing_holiday_exceptions(self):
        dates = ("2026-10-29", "2026-10-30", "2026-10-31", "2026-11-01", "2026-11-02",
                 "2026-12-14", "2026-12-15", "2027-01-05", "2027-01-06", "2027-01-10", "2027-01-11")
        source = self.workbook(pickups=dates, city_rate=10, airport_rate=10)
        _, output, imported = self.run_updates([], source)
        for path in (output, imported):
            for col in range(9, 15):
                for (group, zone, pickup), rate in self.rates(path, col).items():
                    protected = pickup in {"2026-10-31", "2026-11-01", "2026-12-15", "2027-01-05", "2027-01-06", "2027-01-10"}
                    old = 10 + (group in {"EDAV", "EDMV"})
                    expected = old
                    if not protected and group in ALLOWED and col < 14:
                        expected = (60 if zone == "WA1" else 30) + (group in {"EDAV", "EDMV"})
                    if group == "PDAH" and col == 14 and not protected:
                        expected = 170
                    if "2026-10-30" <= pickup <= "2026-11-01":
                        expected = max(expected, 110)
                    if group == "FVMD" and "2026-12-15" <= pickup <= "2027-01-05":
                        expected = (1000, 700, 600, 500, 500, 400)[col - 9]
                    self.assertEqual(rate, expected, (group, zone, pickup, col))

    def test_floor_intersects_configured_groups_and_exclusions(self):
        for override in ({"apply_groups": ["CDMV"]}, {"excluded_groups": FROZEN + ["EDMV"]}):
            with self.subTest(override=override):
                source = self.workbook(city_rate=10, airport_rate=10)
                _, output, _ = self.run_updates([], source, {**self.config, **override})
                prices = self.rates(output, 10)
                self.assertEqual(prices["CDMV", "WA1", "2026-10-01"], 60)
                self.assertEqual(prices["EDMV", "WA1", "2026-10-01"], 11)

    def test_registry_mirrors_not_names_determine_kind(self):
        registry = json.loads(FIXTURE_REGISTRY_PATH.read_text(encoding="utf-8"))
        for location in registry["locations"]:
            location["scraper_label"] = "Airport-looking city" if location["kind"] == "city" else "Downtown-looking airport"
            location["scraper_label"] += location["id"]
        registry["zone_mirrors"]["NOT_AN_AIRPORT"] = "KRTI"
        path = self.root / "registry.json"
        path.write_text(json.dumps(registry), encoding="utf-8")
        path = self.root / "config.json"
        path.write_text(json.dumps({"location_registry_file": "registry.json",
                                    "pricing_rules_file": self.config["pricing_rules_file"]}), encoding="utf-8")
        config = load_config(path)
        for zone, expected in (("WR2", 60), ("KRTI", 30), ("NOT_AN_AIRPORT", 30), ("UNKNOWN_AIRPORT", None)):
            self.assertEqual(get_mandatory_floor(zone, "CDMV", 1, 1, config, "2026-10-13")[0], expected)

    def test_invalid_new_rule_fails_closed_before_outputs_are_saved(self):
        original = load_fixture_pricing_rules()
        for override in ({"groups": ["cdmv"]}, {"groups": ["CDMV", "CDMV"]}, {"groups": ["*"]},
                         {"groups": ["CDMVX"]}, {"maxDays": 21}, {"minDays": True},
                         {"minimumRatesByKind": {"city": 60}},
                         {"minimumRatesByKind": {"city": 60, "airport": 35, "other": 50}},
                         {"minimumRatesByKind": {"city": True, "airport": 35}},
                         {"minimumRatesByKind": {"city": 60, "airport": float("inf")}}):
            rules = json.loads(json.dumps(original))
            rules["pricing"]["locationTypeFloorsPlnDay"].update(override)
            path = self.root / "invalid-new-rule.json"
            path.write_text(json.dumps(rules), encoding="utf-8")
            with self.subTest(override=override), self.assertRaisesRegex(ValueError, "locationTypeFloorsPlnDay"):
                self.run_updates([], self.workbook(), {**self.config, "pricing_rules_file": str(path)})
            self.assertFalse((self.root / "review.xlsx").exists())

    def test_cli_market_selection_does_not_narrow_mandatory_allowlist(self):
        source = self.workbook(city_rate=10, airport_rate=10)
        recommendations = self.root / "empty.json"
        recommendations.write_text('{"decisions": []}', encoding="utf-8")
        summary = apply_updates(source, recommendations, None, self.config, "CDMV", True)
        self.assertEqual({change["group"] for change in summary["changes"] if change["duration_max_days"] <= 20}, set(ALLOWED))

    def test_post_priority_city_floor_conflict_ignores_cap_and_applies_market_target(self):
        source = self.workbook(pickups=("2026-10-26",), city_rate=32, airport_rate=30)
        original = source.read_bytes()
        summary, output, imported = self.run_updates([
            self.decision(pickup="2026-10-26", suggested=79, mm_rank=2)
        ], source)
        self.assertEqual(self.rates(output, 10)["CDMV", "WA1", "2026-10-26"], 79)
        self.assertEqual(summary["city_top1_airport_cap_conflict_count"], 0)
        self.assertEqual(summary["city_top1_airport_cap_conflicts"], [])
        self.assertEqual(summary["skipped_target_count"], 0)
        self.assertEqual(summary["city_top1_airport_cap_ignored_count"], 1)
        scope, = summary["city_top1_airport_cap_ignored_scopes"]
        self.assertEqual((scope["zone"], scope["pickup_date"], scope["duration_band"]), ("WA1", "2026-10-26", "2"))
        self.assertEqual(scope["groups"], sorted(ALLOWED))
        self.assertIn("60", scope["reason"])
        self.assertIn("39", scope["reason"])
        self.assertEqual(summary["city_top1_airport_cap_violation_count"], 0)
        self.assertFalse(any(row["status"] == "FAIL" for row in summary["validation"]))
        market_changes = [change for change in summary["changes"] if change["zone"] == "WA1" and change["duration_band"] == "2"]
        self.assertEqual({change["group"] for change in market_changes}, set(ALLOWED))
        for change in market_changes:
            self.assertEqual(change["target_rank"], 1)
            self.assertTrue(change["city_top1_airport_cap_ignored"])
            self.assertFalse(change["city_top1_airport_cap_active"])
            self.assertFalse(change["city_top1_airport_cap_applied"])
            self.assertEqual(change["city_top1_airport_cap_ignored_reason"], scope["reason"])
        for path in (output, imported):
            prices = self.rates(path, 10)
            for group in ALLOWED:
                self.assertEqual(prices[group, "WA1", "2026-10-26"], 79 + (group in {"EDAV", "EDMV"}))
                self.assertEqual(prices[group, "WALO", "2026-10-26"], 30 + (group in {"EDAV", "EDMV"}))
            for group in FROZEN:
                self.assertEqual(prices[group, "WA1", "2026-10-26"], 32)
        self.assertEqual(source.read_bytes(), original)
        book = openpyxl.load_workbook(output)
        rows = list(book["Recommendations Review"].values)
        headers = rows[0]
        records = [dict(zip(headers, row)) for row in rows[1:]]
        self.assertFalse(any(row["Typ rekomendacji"] == "Konflikt floor i limitu lotniska" for row in records))
        market_rows = [row for row in records if "pominieto limit" in str(row["Uwagi kontroli"]).lower()]
        self.assertTrue(market_rows)
        for row in market_rows:
            self.assertIn("obowiazujace minimum", row["Uwagi kontroli"].lower())
            self.assertIn("60", row["Powod"])
            self.assertIn("39", row["Powod"])
            self.assertTrue(row["Prognoza na stronie"])
        book.close()
        preview = apply_updates(source, self.root / "recommendations.json", None, self.config, None, True)
        for field in ("changes", "change_statistics", "skipped_target_count", "city_top1_airport_cap_ignored_count",
                      "city_top1_airport_cap_ignored_scopes", "city_top1_airport_cap_conflict_count", "city_top1_airport_cap_conflicts"):
            self.assertEqual(preview[field], summary[field], field)
        self.assertEqual(source.read_bytes(), original)

    def test_post_priority_city_floor_compatible_cap_still_clamps_market_target(self):
        source = self.workbook(pickups=("2026-10-26",), city_rate=70, airport_rate=50)
        original = source.read_bytes()
        summary, output, imported = self.run_updates([self.decision(pickup="2026-10-26", suggested=79, mm_rank=2)], source)
        self.assertEqual(summary["city_top1_airport_cap_ignored_count"], 0)
        self.assertEqual(summary["city_top1_airport_cap_ignored_scopes"], [])
        self.assertEqual(summary["city_top1_airport_cap_applied_count"], len(ALLOWED))
        self.assertEqual(summary["city_top1_airport_cap_conflict_count"], 0)
        self.assertEqual(summary["skipped_target_count"], 0)
        for change in summary["changes"]:
            if change["group"] not in ALLOWED:
                continue
            self.assertTrue(change["city_top1_airport_cap_active"])
            self.assertTrue(change["city_top1_airport_cap_applied"])
            self.assertFalse(change["city_top1_airport_cap_ignored"])
            self.assertEqual(change["target_rank"], 1)
        for path in (output, imported):
            prices = self.rates(path, 10)
            for group in ALLOWED:
                self.assertEqual(prices[group, "WA1", "2026-10-26"], 65 + (group in {"EDAV", "EDMV"}))
                self.assertEqual(prices[group, "WALO", "2026-10-26"], 50 + (group in {"EDAV", "EDMV"}))
        self.assertEqual(source.read_bytes(), original)

    def test_ignored_city_cap_still_enforces_floor_for_low_market_recommendation(self):
        source = self.workbook(pickups=("2026-10-26",), city_rate=70, airport_rate=30)
        original = source.read_bytes()
        summary, output, imported = self.run_updates([
            self.decision(pickup="2026-10-26", suggested=25, action="decrease", mm_rank=2)
        ], source)
        self.assertEqual(summary["city_top1_airport_cap_ignored_count"], 1)
        self.assertEqual(summary["skipped_target_count"], 0)
        for path in (output, imported):
            prices = self.rates(path, 10)
            for group in ALLOWED:
                self.assertEqual(prices[group, "WA1", "2026-10-26"], 60 + (group in {"EDAV", "EDMV"}))
        for change in summary["changes"]:
            if change["group"] in ALLOWED:
                self.assertEqual(change["target_rank"], 1)
                self.assertTrue(change["city_top1_airport_cap_ignored"])
                self.assertFalse(change["city_top1_airport_cap_active"])
                self.assertFalse(change["target_achievable"])
        self.assertEqual(source.read_bytes(), original)

    def test_cap_ignore_does_not_bypass_evidence_or_protected_dates(self):
        for extra in ({"incomplete_band": True}, {"data_quality_status": "missing_mm_rate"}, {"pickup": "2027-01-10"}):
            with self.subTest(extra=extra):
                source = self.workbook(pickups=("2026-10-26", "2027-01-10"), city_rate=70, airport_rate=30)
                original = source.read_bytes()
                before = {column: self.rates(source, column) for column in (10, 11, 13)}
                decision = self.decision(pickup="2026-10-26", suggested=79)
                if "pickup" in extra:
                    decision = self.decision(pickup=extra["pickup"], suggested=79)
                else:
                    decision.update(extra)
                decisions = ([self.decision(day, pickup="2026-10-26", suggested=79) for day in range(8, 14)]
                             if extra.get("incomplete_band") else [decision])
                summary, output, imported = self.run_updates(decisions, source)
                self.assertEqual(summary["city_top1_airport_cap_ignored_count"], 0)
                self.assertEqual(summary["city_top1_airport_cap_ignored_scopes"], [])
                self.assertTrue(all(change["target_rank"] == "" for change in summary["changes"]))
                for path in (output, imported):
                    for column in (10, 11, 13):
                        self.assertEqual(self.rates(path, column), before[column])
                self.assertEqual(source.read_bytes(), original)

    def test_cap_ignore_intersects_allowlist_and_never_updates_excluded_classes(self):
        source = self.workbook(pickups=("2026-10-26",), city_rate=70, airport_rate=30)
        original = source.read_bytes()
        before = self.rates(source, 10)
        config = {**self.config, "apply_groups": ["CDMV", "PDAH"]}
        summary, output, imported = self.run_updates([self.decision(pickup="2026-10-26", suggested=79)], source, config)
        scope, = summary["city_top1_airport_cap_ignored_scopes"]
        self.assertEqual(scope["groups"], ["CDMV"])
        self.assertEqual(summary["skipped_target_count"], 0)
        for path in (output, imported):
            prices = self.rates(path, 10)
            for key, rate in before.items():
                self.assertEqual(prices[key], 79 if key == ("CDMV", "WA1", "2026-10-26") else rate)
        self.assertEqual(source.read_bytes(), original)

    def test_actual_js_rank_fail_hold_still_corrects_python_output(self):
        node = shutil.which("node")
        self.assertIsNotNone(node)
        source = self.workbook(city_rate=10, airport_rate=500)
        offers = [{"provider_name": provider, "total_price": rate * 2, "rental_days": 2, "currency": "PLN"}
                  for provider, rate in (("Cheap One", 25), ("Cheap Two", 26), ("Cheap Three", 27), ("MM Cars Rental", 30))]
        view = {"top_3": offers[:3], "mm_cars_rental": offers[3]}
        payload = self.root / "results.json"
        payload.write_text(json.dumps({"locations": ["Warsaw West Train Station"], "scenarios": [{
            "start_date": "2026-10-01", "rental_days": 2,
            "top_3_plus_mm_by_location": {"Warsaw West Train Station": view},
            "offer_views_by_location": {"Warsaw West Train Station": {"automatic": view}},
        }]}), encoding="utf-8")
        recommendations = self.root / "js-recommendations.json"
        subprocess.run([node, "src/pricingRecommendations.js", str(payload), str(recommendations),
                        f"--config={self.config['pricing_rules_file']}"], cwd=ROOT, check=True, capture_output=True)
        decisions = json.loads(recommendations.read_text(encoding="utf-8"))["decisions"]
        self.assertTrue(decisions)
        self.assertTrue(all(item["action"] == "hold" for item in decisions))
        self.assertTrue(any(item["data_quality_status"] == "floor_blocks_top3" for item in decisions), decisions)
        summary, output, _ = self.run_updates(decisions, source)
        self.assertEqual(self.rates(output, 10)["CDMV", "WA1", "2026-10-01"], 60)
        self.assertEqual(self.rates(output, 10)["EDAV", "WA1", "2026-10-01"], 61)
        self.assertTrue(all(change["target_rank"] == "" for change in summary["changes"]))

    def test_dry_run_matches_export_and_validation_reads_final_floor_cells(self):
        source = self.workbook(city_rate=10, airport_rate=10)
        recommendations = self.root / "empty.json"
        recommendations.write_text('{"decisions": []}', encoding="utf-8")
        preview = apply_updates(source, recommendations, None, self.config, None, True)
        summary, output, _ = self.run_updates([], source)
        self.assertEqual(preview["changes"], summary["changes"])
        self.assertEqual(summary["change_count"], 62)  # 2 zones * 6 groups * 5 bands, plus PDAH long minima.
        book = openpyxl.load_workbook(output)
        sheet = book["Sheet1"]
        sheet["J5"] = 59
        checks = build_validation_rows(sheet, self.config, get_duration_columns(sheet, self.config), [], [])
        self.assertEqual(next(c[2] for c in checks if c[0] == "Stawki ponizej obowiazkowego minimum typu lokalizacji"), 1)
        book.close()

    def test_clean_import_preserves_merged_headers_and_handles_frozen_split(self):
        source = self.workbook()
        book = openpyxl.load_workbook(source)
        sheet = book["Sheet1"]
        sheet.merge_cells("A1:H1")
        sheet.sheet_view.pane.state = "frozenSplit"
        book.save(source)
        book.close()
        original = source.read_bytes()
        _, output, imported = self.run_updates([], source)
        for path in (output, imported):
            book = openpyxl.load_workbook(path)
            self.assertEqual([str(value) for value in book["Sheet1"].merged_cells.ranges], ["A1:H1"])
            self.assertTrue(all(sheet.freeze_panes is None for sheet in book))
            if path == imported:
                self.assertTrue(all(cell.comment is None for row in book["Sheet1"] for cell in row))
            book.close()
        self.assertEqual(source.read_bytes(), original)

    def test_priority_legend_does_not_present_superseded_30_and_40_floors_as_active(self):
        _, output, _ = self.run_updates([], self.workbook())
        book = openpyxl.load_workbook(output)
        rows = list(book['Changed Positions'].values)
        text = next(str(row[1]) for row in rows if row[0] == 'Przebicie top1')
        self.assertIn('miasto 60 PLN', text)
        self.assertIn('lotnisko 30 PLN', text)
        self.assertNotIn('floor 30 PLN', text)
        self.assertNotIn('floor 40 PLN', text)
        book.close()


if __name__ == "__main__":
    unittest.main(verbosity=2)
