import sys
from datetime import date, timedelta
from pathlib import Path

import openpyxl

target = Path(sys.argv[1])
count = int(sys.argv[2])
book = openpyxl.Workbook()
sheet = book.active
sheet.title = "Sheet1"
sheet.append(["Rental rates for packages: INCLUSIVE POA"])
sheet.append(["Min days"] + [None] * 7 + [2])
sheet.append(["Max days"] + [None] * 7 + [2])
sheet.append(["Group", "Description", "Rate code", "Zone", "Booking start date",
              "Booking end date", "Pickup start date", "Pickup end date", "Per day"])
for offset in range(count):
    pickup = date(2026, 10, 16) + timedelta(days=offset)
    sheet.append(["CDMV", None, "POA", "TEST", date(2026, 10, 1), pickup, pickup, pickup, 80])
if len(sys.argv) > 3 and sys.argv[3] == "missing-band":
    sheet.cell(2, 10).value = 5
    sheet.cell(3, 10).value = 7
    sheet.cell(4, 10).value = "Per day"
book.save(target)
book.close()
