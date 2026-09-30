import hashlib
import json
import shutil
import subprocess
import sys
import tempfile
from datetime import date
from pathlib import Path

import openpyxl
from openpyxl.styles import PatternFill

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from tools.update_excel_rates import (apply_updates, load_config, predict_site_rate,
    evaluate_target_constraints, target_for_group, format_broker_markup)


def main():
    fixed = json.loads((ROOT / 'input/broker-markup-frozen.json').read_text())['brokerMarkupCalibration']
    groups = fixed['baseGroups'] + list(fixed['groupSupplementsPlnDay'])
    node = shutil.which('node')
    assert node, 'Node is required for the end-to-end recommendation test'
    with tempfile.TemporaryDirectory(prefix='fixed-markup-excel-') as directory:
        folder = Path(directory)
        source = folder / 'baseline.xlsx'
        workbook = openpyxl.Workbook()
        ws = workbook.active
        ws.title = 'Sheet1'
        ws.append(['Rental rates for packages: INCLUSIVE POA'])
        ws.append(['Min days'] + [None]*7 + [1,2,3,5,8,21])
        ws.append(['Max days'] + [None]*7 + [1,2,4,7,20,35])
        ws.append(['Group','Description','Rate code','Zone','Booking start date', 'Booking end date',
            'Pickup start date','Pickup end date'] + ['Per day']*6)
        ws['A1'].fill = PatternFill(fill_type='solid', fgColor='1F4E78')
        for start in [date(2026,10,1), date(2026,10,31)]:
            for group in groups:
                ws.append([group,None,None,'WALO',start,start,start,start]+[150]*6)
        ws.auto_filter.ref = f'A4:N{ws.max_row}'
        ws.row_dimensions[5].hidden = True
        workbook.save(source)
        saved_source = openpyxl.load_workbook(source)
        headers = [[(c.value, str(c._style)) for c in row] for row in saved_source['Sheet1'].iter_rows(max_row=4)]
        saved_source.close()
        original_hash = hashlib.sha256(source.read_bytes()).hexdigest()
        scenarios=[]
        for start in ['2026-10-01','2026-10-31']:
            for days in range(2,15):
                def offer(provider, rate):
                    return dict(provider_name=provider,total_price=rate*days,rental_days=days,currency='PLN')
                view={'top_3':[offer('Competitor',100),offer('MM Cars Rental',120),offer('Other',140)],
                    'mm_cars_rental':offer('MM Cars Rental',120)}
                scenarios.append(dict(start_date=start,rental_days=days,
                    top_3_plus_mm_by_location={'Warsaw Chopin Airport (WAW)':view},
                    offer_views_by_location={'Warsaw Chopin Airport (WAW)':{'automatic':view}}))
        payload = folder / 'results.json'
        payload.write_text(json.dumps({'locations':['Warsaw Chopin Airport (WAW)'],'scenarios':scenarios}))
        recommendations = folder / 'recommendations.json'
        wrong = folder / 'wrong-calibration.json'
        wrong.write_text(json.dumps({'brokerMarkupCalibration':{'enabled':True,'defaultMultiplier':10}}))
        subprocess.run([node,'src/pricingRecommendations.js',str(payload),str(recommendations),
            '--config=pricing-rules.config.example.json',f'--calibration={wrong}'],cwd=ROOT,check=True,capture_output=True)
        decisions=json.loads(recommendations.read_text())
        assert decisions['decisions'][0]['suggested_rate_pln_day'] == 70
        assert decisions['decisions'][0]['predicted_site_rate_pln_day'] == 99
        config=load_config(ROOT / 'excel-rate-update.config.example.json')
        config['pickup_date_expansion']={'enabled':False}
        config['baseline_manifest_file']=''
        config['baseline_workbook_sha256']=original_hash
        report=folder / 'recommendations.xlsx'
        imported=folder / 'import.xlsx'
        summary=apply_updates(source,recommendations,report,config,None,False,import_output_path=imported)
        assert not summary['broker_markup_observations']['enabled']
        assert summary['broker_markup_observations']['count']==0
        assert summary['change_count']==18, summary['change_count']
        assert hashlib.sha256(source.read_bytes()).hexdigest()==original_hash
        result=openpyxl.load_workbook(imported)
        assert result.sheetnames==['Sheet1']
        sheet=result['Sheet1']
        assert sheet.max_row==ws.max_row
        for before, after in zip(headers, sheet.iter_rows(max_row=4)):
            assert [(c.value,str(c._style)) for c in after]==before, (before,[(c.value,str(c._style)) for c in after])
        active={'CDMV','CGAV','CWAV','CWMR','EDAV','EDMV'}
        for row in sheet.iter_rows(min_row=5):
            group=row[0].value
            if row[6].value.date()==date(2026,10,1) and group in active:
                adjustment=1 if group=='EDMV' else 0
                assert [c.value for c in row[9:12]]==[69+adjustment,84+adjustment,90+adjustment]
            else:
                assert all(c.value==150 for c in row[9:12])
            assert [row[i].value for i in [8,12,13]]==[150,150,150]
        for change in summary['changes']:
            assert change['predicted_site_rate']==change['new_rate']+change['broker_markup_amount_pln_day']
            assert change['target_achievable']
        report_book=openpyxl.load_workbook(report)
        assert any('PLN/doba (staly)' in str(c.value) for row in report_book['Recommendations Review'] for c in row)
        assert any('Staly narzut brokera' == c.value for row in report_book['Changed Positions'] for c in row)
        result.close()
        report_book.close()
        workbook.close()

    target={'broker_markup_model':'fixed_amount','broker_markup_multiplier':1,'broker_markup_amount_pln_day':29,
        'broker_markup_group_supplements_pln_day':{'PDAH':10},'suggested_rate_pln_day':70,'site_cap_rate_pln_day':99}
    premium=target_for_group(target,'PDAH')
    assert premium['suggested_rate_pln_day']==60
    assert predict_site_rate(60,premium)==99
    assert predict_site_rate(70,target)==99
    assert evaluate_target_constraints(target,70)['target_achievable']
    assert not evaluate_target_constraints(target,71)['target_achievable']
    assert format_broker_markup(premium)=='39 PLN/doba (staly)'
    print('PASS fixed markup Excel end-to-end: conversion, frozen CLI, exclusions, protected dates, duration8+, headers and baseline preserved')


if __name__=='__main__':
    main()
