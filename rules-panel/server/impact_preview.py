"""Isolated real-export bridge. Called only with a staged preview task."""

import json
import math
import sys
from datetime import date, datetime
from pathlib import Path

sys.dont_write_bytecode = True
ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))

import openpyxl
from tools.update_excel_rates import (
    apply_updates, get_duration_columns, load_config, parse_date_value,
)


class PreviewBlocked(ValueError):
    def __init__(self, code, message):
        super().__init__(message)
        self.code = code


def guard(condition, code, message):
    if not condition:
        raise PreviewBlocked(code, message)


def write_json(path, value):
    path.write_text(json.dumps(value, ensure_ascii=False), encoding="utf-8")


def normalized(value):
    return value.isoformat() if isinstance(value, (date, datetime)) else value


def grid(path, config):
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=False)
    try:
        guard(workbook.sheetnames == ["Sheet1"], "invalid_import", "Clean export must contain only Sheet1.")
        sheet = workbook["Sheet1"]
        columns = sorted(set(get_duration_columns(sheet, config).values()))
        guard(bool(columns), "invalid_import", "Export has no duration columns.")
        output = {}
        identities = set()
        for row in sheet.iter_rows(min_row=5):
            pickup = parse_date_value(row[6].value)
            end = parse_date_value(row[7].value)
            guard(pickup is not None and end is not None, "invalid_import", "Invalid exported pickup dates.")
            identity = [normalized(cell.value) for cell in row[:8]]
            identity_key = json.dumps(identity, ensure_ascii=False)
            guard(identity_key not in identities, "ambiguous_grid", "Duplicate semantic tariff rows cannot be compared safely.")
            identities.add(identity_key)
            for column, band, low, high in columns:
                cell = row[column - 1]
                rate = cell.value
                guard(rate is None or (type(rate) in (int, float) and math.isfinite(rate) and rate >= 0),
                      "invalid_import", "Formula, nonnumeric or invalid exported rate encountered.")
                position_key = json.dumps([identity, low, high], ensure_ascii=False)
                output[position_key] = {"rate": rate, "cell": cell.coordinate, "group": row[0].value,
                                        "zone": row[3].value, "pickupDate": pickup.isoformat(),
                                        "pickupEndDate": end.isoformat(), "durationBand": band}
        return output, sheet.max_row
    finally:
        workbook.close()


def review_decisions(path, sheet_name):
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=False)
    try:
        output = {}
        if sheet_name not in workbook.sheetnames:
            return output
        rows = workbook[sheet_name].iter_rows(values_only=True)
        headers = list(next(rows))
        for row in rows:
            record = dict(zip(headers, row))
            for coordinate in str(record.get("Komorki Excel") or "").split(", "):
                if coordinate:
                    output[coordinate] = {"status": record.get("Status"), "reason": record.get("Powod"),
                                          "controlNotes": record.get("Uwagi kontroli"),
                                          "observedTargetOutcome": record.get("Cel zmiany")}
        return output
    finally:
        workbook.close()


def export(task, label, directory):
    documents = task["current" if label == "before" else "proposed"]
    excel = dict(documents["excel"])
    excel.update({"baseline_manifest_file": task["artifacts"]["baselineManifest"],
                  "pricing_rules_file": str(directory / f"{label}-pricing.json"),
                  "location_registry_file": str(directory / "locations.json")})
    write_json(directory / f"{label}-pricing.json", documents["pricing"])
    write_json(directory / f"{label}-excel.json", excel)
    write_json(directory / f"{label}-recommendations.json", task[label])
    config = load_config(directory / f"{label}-excel.json")
    report = directory / f"{label}-report.xlsx"
    imported = directory / f"{label}-import.xlsx"
    summary = apply_updates(Path(task["artifacts"]["baselineWorkbook"]), directory / f"{label}-recommendations.json",
                            report, config, None, False, import_output_path=imported)
    guard(all(item["status"] not in {"ERROR", "FAIL"} for item in summary["validation"]),
          "validation_failed", "Real export validation reported an error.")
    rates, rows = grid(imported, config)
    return {"changeCount": summary["change_count"], "changeStatistics": summary["change_statistics"],
            "activeRecommendationCount": sum(item.get("action") != "hold" for item in task[label]["decisions"]),
            "holdDecisionCount": sum(item.get("action") == "hold" for item in task[label]["decisions"]),
            "importRowCount": rows, "rateCellCount": len(rates), "skippedTargetCount": summary["skipped_target_count"],
            "validation": summary["validation"], "recommendationDateScope": summary["recommendation_date_scope"],
            "floorChangeCount": summary["mandatory_zone_floor_change_count"] + summary["mandatory_group_floor_change_count"],
            "cityCapConflictCount": summary["city_top1_airport_cap_conflict_count"]}, rates, review_decisions(report, config["recommendations_review_sheet"])


def run(task_path):
    task = json.loads(task_path.read_text(encoding="utf-8"))
    directory = task_path.parent
    write_json(directory / "locations.json", task["current"]["locations"])
    before, before_rates, before_reasons = export(task, "before", directory)
    after, after_rates, after_reasons = export(task, "after", directory)
    changed = increased = decreased = added = removed = 0
    samples = []
    preserved = {"status": "baseline_preserved", "reason": "Silnik nie zmienil stawki w tej pozycji."}
    for key in sorted(before_rates.keys() | after_rates.keys()):
        left = before_rates.get(key)
        right = after_rates.get(key)
        if left is not None and right is not None and left["rate"] == right["rate"]:
            continue
        changed += 1
        if left is None:
            added += 1
        elif right is None:
            removed += 1
        elif left["rate"] is not None and right["rate"] is not None:
            increased += right["rate"] > left["rate"]
            decreased += right["rate"] < left["rate"]
        if len(samples) < 20:
            position = right or left
            samples.append({"position": {name: value for name, value in position.items() if name != "rate"},
                            "beforeCell": left["cell"] if left else None, "afterCell": right["cell"] if right else None,
                            "beforeRate": left["rate"] if left else None, "afterRate": right["rate"] if right else None,
                            "beforeDecision": before_reasons.get(left["cell"], preserved) if left else {"status": "absent"},
                            "afterDecision": after_reasons.get(right["cell"], preserved) if right else {"status": "absent"}})
    return {"before": before, "after": after, "impact": {"changedRateCells": changed,
            "increasedRateCells": increased, "decreasedRateCells": decreased, "addedRateCells": added,
            "removedRateCells": removed, "unchangedRateCells": len(before_rates.keys() & after_rates.keys()) - (changed - added - removed),
            "samples": samples, "samplesTruncated": changed > len(samples)},
            "limitations": ["Offline replay of recorded evidence, not a live ranking or revenue guarantee."]}


if __name__ == "__main__":
    try:
        result = run(Path(sys.argv[1]))
    except PreviewBlocked as error:
        result = {"error": {"code": error.code, "message": str(error)}}
    except Exception:
        result = {"error": {"code": "engine_failed", "message": "The real export engine rejected the evidence or configuration."}}
    print(json.dumps(result, ensure_ascii=False, allow_nan=False))
